#!/usr/bin/env python3
"""Native OOXML PPTX import/export for supported editable text and image widgets.

The parser does not execute embedded content. Unsupported native features (animations,
SmartArt, complex charts and masters) are reported separately rather than silently
claimed as editable.
"""
import argparse
import io
import os
import stat
import tempfile
import xml.etree.ElementTree as ET
import json
import math
import hashlib
import uuid
from datetime import date
from pathlib import Path
from zipfile import ZipFile, BadZipFile, ZIP_STORED, ZIP_DEFLATED
import re
import posixpath
from urllib.parse import unquote

MAX_PPTX_BYTES = 50 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 250 * 1024 * 1024
MAX_ZIP_ENTRIES = 3000
MAX_ZIP_MEMBER_BYTES = 80 * 1024 * 1024
MAX_ZIP_COMPRESSION_RATIO = 1000
MAX_IMAGE_PIXELS = 16_000_000
MAX_EXPORT_TEXT = 6000


def _supported_image(blob, mime):
    """Validate the actual image stream, not just a spoofable file signature.

    Pillow's verify() checks stream structure; load() checks decompression and
    truncation. An explicit pixel budget prevents tiny compressed image bombs.
    """
    if not isinstance(blob, bytes) or not 0 < len(blob) <= 10 * 1024 * 1024:
        return False
    signatures = {'image/png': b'\x89PNG\r\n\x1a\n', 'image/jpeg': b'\xff\xd8\xff'}
    if mime not in signatures or not blob.startswith(signatures[mime]):
        return False
    try:
        from PIL import Image, UnidentifiedImageError
        with Image.open(io.BytesIO(blob)) as picture:
            if picture.format != {'image/png': 'PNG', 'image/jpeg': 'JPEG'}[mime]:
                return False
            if picture.width < 1 or picture.height < 1 or picture.width * picture.height > MAX_IMAGE_PIXELS:
                return False
            picture.verify()
        with Image.open(io.BytesIO(blob)) as picture:
            picture.load()
        return True
    except (OSError, ValueError, SyntaxError, UnidentifiedImageError, Image.DecompressionBombError):
        return False


def _valid_text(value, maximum=MAX_EXPORT_TEXT):
    """Reject controls prohibited by XML 1.0 instead of creating corrupt OOXML."""
    return (isinstance(value, str) and len(value) <= maximum and
            all(ord(c) in (9, 10, 13) or 32 <= ord(c) <= 0xD7FF or
                0xE000 <= ord(c) <= 0xFFFD or 0x10000 <= ord(c) <= 0x10FFFF
                for c in value))


def _validate_export_input(data):
    if not isinstance(data, dict) or data.get('format') != 'quorivio-session' or data.get('version') != 1:
        raise ValueError('Not a Quorivio session definition')
    activities = data.get('activities')
    if not isinstance(activities, list) or not 1 <= len(activities) <= 200:
        raise ValueError('Invalid page count')
    for activity in activities:
        if not isinstance(activity, dict):
            raise ValueError('Invalid slide')
        design = activity.get('page_design', {})
        if not isinstance(design, dict):
            raise ValueError('Invalid page design')
        background = design.get('backgroundColor')
        if background is not None and (not isinstance(background, str) or not re.fullmatch(r'#[0-9a-fA-F]{6}', background)):
            raise ValueError('Invalid slide background color')
        widgets = design.get('widgets', [])
        if not isinstance(widgets, list) or len(widgets) > 16:
            raise ValueError('Too many or invalid widgets')
        for widget in widgets:
            if not isinstance(widget, dict):
                raise ValueError('Invalid widget')
            kind = widget.get('kind')
            if kind not in ('text', 'postit', 'question', 'timed-quiz', 'image', 'logo'):
                continue  # Other widget kinds have no native PowerPoint representation.
            position = widget.get('position', {'x': 5, 'y': 5, 'width': 90, 'height': 20})
            if not isinstance(position, dict):
                raise ValueError('Invalid widget geometry')
            raw_coordinates = [position.get(k, v) for k, v in
                               [('x', 5), ('y', 5), ('width', 90), ('height', 20)]]
            if any(isinstance(v, bool) or not isinstance(v, (int, float)) for v in raw_coordinates):
                raise ValueError('Invalid widget geometry')
            x, y, w, h = raw_coordinates
            if not all(math.isfinite(v) for v in (x, y, w, h)) or x < 0 or y < 0 or w <= 0 or h <= 0 or x + w > 100.00001 or y + h > 100.00001:
                raise ValueError('Invalid widget geometry')
            rotation = position.get('rotation', 0)
            if isinstance(rotation, bool) or not isinstance(rotation, (int, float)) or not math.isfinite(rotation) or abs(rotation) > 360_000:
                raise ValueError('Invalid widget rotation')
            if kind in ('image', 'logo'):
                asset_id = widget.get('assetId')
                try:
                    valid_asset_id = (isinstance(asset_id, str) and
                                      str(uuid.UUID(asset_id)) == asset_id and
                                      uuid.UUID(asset_id).version == 4)
                except (ValueError, AttributeError, TypeError):
                    valid_asset_id = False
                if not valid_asset_id:
                    raise ValueError('Invalid image asset ID')
                if 'crop' in widget and not _valid_crop(widget['crop']):
                    raise ValueError('Invalid picture crop')
                continue
            if 'crop' in widget:
                raise ValueError('Crop is only supported for pictures')
            for key in ('label', 'description'):
                value = widget.get(key)
                if value is not None and not _valid_text(value):
                    raise ValueError('Invalid widget text')
            style = widget.get('style', {})
            if not isinstance(style, dict):
                raise ValueError('Invalid widget style')
            for key in ('fontFamily', 'color', 'fillColor', 'strokeColor', 'shapeType'):
                if key in style and not isinstance(style[key], str):
                    raise ValueError('Invalid widget style')
            if 'fontFamily' in style and (not _valid_text(style['fontFamily'], 80) or
                                          not style['fontFamily'].strip()):
                raise ValueError('Invalid widget font family')
            for key in ('color', 'fillColor', 'strokeColor'):
                if key in style and not re.fullmatch(r'#[0-9a-fA-F]{6}', style[key]):
                    raise ValueError('Invalid widget color')
            if 'textAlign' in style and style['textAlign'] not in ('left', 'center', 'right'):
                raise ValueError('Invalid text alignment')
            if 'verticalAlign' in style and style['verticalAlign'] not in ('top', 'middle', 'bottom'):
                raise ValueError('Invalid vertical alignment')
            if 'shapeType' in style and style['shapeType'] not in ('ellipse', 'triangle', 'rectangle'):
                raise ValueError('Invalid widget shape')
            for key in ('bold', 'italic', 'underline', 'rounded', 'fillNone', 'strokeNone'):
                if key in style and not isinstance(style[key], bool):
                    raise ValueError('Invalid widget style')
            if 'strokeDash' in style and style['strokeDash'] not in ('solid', 'dash', 'dot'):
                raise ValueError('Invalid shape outline dash')
            if 'strokeWidth' in style and (type(style['strokeWidth']) not in (int, float) or
                not math.isfinite(style['strokeWidth']) or not 0.25 <= style['strokeWidth'] <= 12):
                raise ValueError('Invalid shape outline width')
            if 'fontSize' in style and (isinstance(style['fontSize'], bool) or
                not isinstance(style['fontSize'], (int, float)) or not math.isfinite(style['fontSize']) or
                not 1 <= style['fontSize'] <= 200):
                raise ValueError('Invalid widget style')
            if kind in ('question', 'timed-quiz'):
                options = widget.get('options', [])
                if not isinstance(options, list) or len(options) > 12 or any(
                    not isinstance(option, dict) or not _valid_text(option.get('label', ''), 500)
                    for option in options):
                    raise ValueError('Invalid quiz options')
    return activities


SUPPORTED_FONTS = {"Arial", "Georgia", "Verdana", "Tahoma", "Trebuchet MS", "monospace"}

def validate_pptx_archive(source):
    """Reject oversized/bomb-like OOXML packages before python-pptx parses them."""
    path = Path(source)
    if path.is_symlink() or not path.is_file() or not stat.S_ISREG(path.stat().st_mode):
        raise ValueError('PPTX must be a regular file')
    if path.stat().st_size > MAX_PPTX_BYTES:
        raise ValueError("PPTX exceeds 50 MiB import limit")
    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES or sum(e.file_size for e in entries) > MAX_UNCOMPRESSED_BYTES:
                raise ValueError("PPTX decompression limits exceeded")
            if any(e.compress_type not in (ZIP_STORED, ZIP_DEFLATED) for e in entries):
                raise ValueError('Unsupported OOXML compression method')
            if any(e.file_size > MAX_ZIP_MEMBER_BYTES or
                   (e.file_size > 1024 * 1024 and e.file_size > max(e.compress_size, 1) * MAX_ZIP_COMPRESSION_RATIO)
                   for e in entries):
                raise ValueError('Suspicious OOXML compression ratio or member size')
            if any(((e.external_attr >> 16) & 0o170000) == 0o120000 for e in entries):
                raise ValueError('OOXML symbolic links are forbidden')
            if any(e.filename.startswith('/') or '..' in e.filename.split('/') or '.' in e.filename.split('/') or '\\' in e.filename or ':' in e.filename or '\x00' in e.filename or e.filename.startswith('./') or '//' in e.filename or any(ord(c) < 32 for c in e.filename) or (e.flag_bits & 1) for e in entries):
                raise ValueError('Unsafe OOXML archive member')
            if len({e.filename.casefold() for e in entries}) != len(entries):
                raise ValueError('Duplicate OOXML archive member')
            if any(e.filename.lower().endswith(("vbaproject.bin", ".exe", ".dll", ".vbs", ".js", ".ps1", ".bat", ".cmd", ".com", ".scr", ".jar")) for e in entries):
                raise ValueError("Macro or executable content is not supported")
            if any(e.filename.casefold().startswith('ppt/activex/') or
                   (e.filename.casefold().startswith('ppt/embeddings/') and e.filename.casefold().endswith('.bin'))
                   for e in entries):
                raise ValueError('Embedded active OOXML objects are not supported')
            if not {'[Content_Types].xml', 'ppt/presentation.xml'}.issubset({e.filename for e in entries}):
                raise ValueError("Not an OOXML presentation")
            # XML DTD/entity declarations can create parser expansion attacks.
            # Scan bounded XML parts before python-pptx/libreoffice sees them.
            for entry in entries:
                if entry.filename.lower().endswith(('.xml', '.rels')) or entry.filename == '[Content_Types].xml':
                    if entry.file_size > 20 * 1024 * 1024:
                        raise ValueError('Oversized OOXML XML part')
                    with archive.open(entry) as part:
                        # Scan the entire XML part in bounded chunks: declarations
                        # placed after a large comment must not bypass preflight.
                        tail = b''
                        while block := part.read(64 * 1024):
                            segment = (tail + block).lower()
                            if b'<!doctype' in segment or b'<!entity' in segment:
                                raise ValueError('OOXML DTD/entities are not supported')
                            tail = segment[-16:]
            # A renamed PPTM must not bypass the extension-based macro policy.
            with archive.open('[Content_Types].xml') as types_part:
                types_prefix = types_part.read(20 * 1024 * 1024).lower()
                if any(marker in types_prefix for marker in
                       (b'macroenabled', b'vbaproject', b'activex', b'oleobject')):
                    raise ValueError('Active OOXML content type is not supported')
            # OOXML relationship parts can reference external resources; never
            # let a document import silently carry active external references.
            for entry in entries:
                if entry.filename.casefold().endswith('.rels'):
                    if entry.file_size > 2 * 1024 * 1024:
                        raise ValueError('Oversized OOXML relationships')
                    payload = archive.read(entry)
                    if re.search(rb'TargetMode\s*=\s*["\']External["\']', payload, re.I):
                        raise ValueError('External OOXML relationships are not supported')
                    # XML parsing catches single/double quoted, namespace and
                    # entity-normalized attributes that regex-only scanning misses.
                    try:
                        root = ET.fromstring(payload)
                    except ET.ParseError as exc:
                        raise ValueError('Malformed OOXML relationships') from exc
                    for relation in root.iter():
                        if relation.tag.rsplit('}', 1)[-1].lower() != 'relationship':
                            continue
                        attributes = {key.rsplit('}', 1)[-1].lower(): value.strip()
                                      for key, value in relation.attrib.items()}
                        if attributes.get('targetmode', '').lower() == 'external':
                            raise ValueError('External OOXML relationships are not supported')
                        rel_type = attributes.get('type', '').lower()
                        if any(marker in rel_type for marker in ('/oleobject', '/activexcontrol', '/vbaproject', '/activexcontrolbinary')):
                            raise ValueError('Active OOXML relationship is not supported')
                        target = attributes.get('target', '')
                        if not target or any(ord(c) < 32 or ord(c) == 127 for c in target):
                            raise ValueError('Unsafe OOXML relationship target')
                        # Normalize percent-encoded URI characters, including a
                        # second encoded layer; malformed escape sequences fail.
                        if re.search(r'%(?![0-9a-fA-F]{2})|%(?:2f|5c|3a)', target, re.I):
                            raise ValueError('Unsafe OOXML relationship target')
                        for _ in range(8):
                            decoded = unquote(target)
                            if decoded == target:
                                break
                            target = decoded
                        if '%' in target and re.search(r'%[0-9a-fA-F]{2}', target):
                            raise ValueError('Unsafe OOXML nested percent encoding')
                        # Internal targets may legitimately contain ../ to refer
                        # to other package parts, but never a URI or absolute path.
                        if (target.startswith(('/', '\\')) or '\\' in target or
                                re.match(r'^[a-z][a-z0-9+.-]*:', target, re.I) or
                                re.search(r'%(?:2f|5c|3a)', target, re.I) or
                                any(ord(c) < 32 or ord(c) == 127 for c in target)):
                            raise ValueError('Unsafe OOXML relationship target')
                        # Resolve relative to the owner part, not the _rels directory.
                        # ../media/foo is legitimate; traversal beyond package root is not.
                        rel_dir = posixpath.dirname(entry.filename)
                        owner_dir = posixpath.dirname(rel_dir)
                        if rel_dir == '_rels':
                            owner_dir = ''
                        clean_target = target.split('#', 1)[0].split('?', 1)[0]
                        if clean_target and (posixpath.normpath(posixpath.join(owner_dir, clean_target))
                                             in ('..', '.') or
                                             posixpath.normpath(posixpath.join(owner_dir, clean_target)).startswith('../')):
                            raise ValueError('OOXML relationship escapes package root')

    except BadZipFile as exc:
        raise ValueError("Invalid PPTX archive") from exc



def _position(shape, width, height):
    """Preserve small objects instead of enlarging logos/text to legacy 12%/8% minima.

    Off-canvas elements are clipped because the editable Quorivio canvas cannot
    represent them without invalid coordinates. Report that loss separately.
    """
    x = max(0.0, min(99.0, shape.left / width * 100))
    y = max(0.0, min(99.0, shape.top / height * 100))
    w = max(0.01, min(100.0 - x, shape.width / width * 100))
    h = max(0.01, min(100.0 - y, shape.height / height * 100))
    return {'x': round(x, 4), 'y': round(y, 4),
            'width': round(w, 4), 'height': round(h, 4)}


class _PositionedShape:
    """Read-only proxy for a child positioned inside one or more PPTX groups."""
    def __init__(self, original, left, top, width, height, rotation):
        self._original = original
        self.left, self.top = left, top
        self.width, self.height = width, height
        self.rotation = rotation

    def __getattr__(self, name):
        return getattr(self._original, name)


def _flatten_shapes(shapes, transform=None, parent_rotation=0.0, depth=0):
    """Flatten nested OOXML groups while preserving their child coordinate systems.

    This preserves text and supported primitives; the grouping relation, rotations
    involving non-right-angle group transforms and effects are inherently lossy.
    A depth limit rejects pathological nesting instead of exhausting recursion.
    """
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    if depth > 12:
        raise ValueError('PowerPoint groups exceed maximum nesting depth')
    for shape in shapes:
        if transform is None:
            left, top, width, height = shape.left, shape.top, shape.width, shape.height
        else:
            ox, oy, sx, sy = transform
            left = int(round(ox + shape.left * sx))
            top = int(round(oy + shape.top * sy))
            width = int(round(shape.width * sx))
            height = int(round(shape.height * sy))
        rotation = (parent_rotation + float(shape.rotation or 0)) % 360
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            child = shape.shapes
            if not shape.width or not shape.height or not shape._element.grpSpPr.xfrm.chExt.cx or not shape._element.grpSpPr.xfrm.chExt.cy:
                continue
            child_transform = (left - shape._element.grpSpPr.xfrm.chOff.x * width / shape._element.grpSpPr.xfrm.chExt.cx,
                               top - shape._element.grpSpPr.xfrm.chOff.y * height / shape._element.grpSpPr.xfrm.chExt.cy,
                               width / shape._element.grpSpPr.xfrm.chExt.cx, height / shape._element.grpSpPr.xfrm.chExt.cy)
            yield from _flatten_shapes(child, child_transform, rotation, depth + 1)
        else:
            yield _PositionedShape(shape, left, top, width, height, rotation)


def _slide_background(slide):
    """Read only explicit solid RGB fills; never invent a theme color."""
    try:
        from pptx.enum.dml import MSO_FILL
        fill = slide.background.fill
        if fill.type == MSO_FILL.SOLID and fill.fore_color.rgb is not None:
            return '#' + str(fill.fore_color.rgb)
    except (AttributeError, TypeError, ValueError):
        pass
    return None


def _text_frame_style(shape):
    """Extract only consistent paragraph alignment and supported anchoring."""
    style = {}
    if not shape.has_text_frame:
        return style
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    paragraphs = [p for p in shape.text_frame.paragraphs if p.text.strip()]
    aligns = {p.alignment for p in paragraphs}
    if len(aligns) == 1:
        align = next(iter(aligns))
        if align in (PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.RIGHT):
            style['textAlign'] = {PP_ALIGN.LEFT:'left', PP_ALIGN.CENTER:'center', PP_ALIGN.RIGHT:'right'}[align]
    anchor = shape.text_frame.vertical_anchor
    if anchor in (MSO_ANCHOR.TOP, MSO_ANCHOR.MIDDLE, MSO_ANCHOR.BOTTOM):
        style['verticalAlign'] = {MSO_ANCHOR.TOP:'top', MSO_ANCHOR.MIDDLE:'middle', MSO_ANCHOR.BOTTOM:'bottom'}[anchor]
    return style


def _native_shape_paint(shape):
    """Preserve only explicit editable shape paints, avoiding invented theme colors."""
    from pptx.enum.dml import MSO_FILL, MSO_LINE_DASH_STYLE, MSO_COLOR_TYPE
    style = {}
    try:
        if shape.fill.type == MSO_FILL.BACKGROUND:
            style['fillNone'] = True
        elif shape.fill.type == MSO_FILL.SOLID and shape.fill.fore_color.type == MSO_COLOR_TYPE.RGB:
            style['fillColor'] = '#' + str(shape.fill.fore_color.rgb)
    except (AttributeError, TypeError, ValueError):
        pass
    try:
        line = shape.line
        if line.fill.type == MSO_FILL.BACKGROUND:
            style['strokeNone'] = True
        elif line.color.type == MSO_COLOR_TYPE.RGB:
            style['strokeColor'] = '#' + str(line.color.rgb)
            if line.width is not None:
                width = round(line.width.pt, 3)
                if 0.25 <= width <= 12:
                    style['strokeWidth'] = width
            dash = {MSO_LINE_DASH_STYLE.SOLID: 'solid',
                    MSO_LINE_DASH_STYLE.DASH: 'dash',
                    MSO_LINE_DASH_STYLE.ROUND_DOT: 'dot'}.get(line.dash_style)
            if dash:
                style['strokeDash'] = dash
    except (AttributeError, TypeError, ValueError):
        pass
    return style


def _picture_crop(shape):
    """Recover supported PowerPoint picture cropping as four normalized fractions.

    Native OOXML permits negative crop/overscan; those cases are explicitly
    lossy rather than silently introducing invalid browser geometry.
    """
    values = [float(getattr(shape, 'crop_' + side, 0) or 0)
              for side in ('left', 'right', 'top', 'bottom')]
    if not all(math.isfinite(value) and 0 <= value <= 0.95 for value in values):
        return None
    if values[0] + values[1] >= 0.99 or values[2] + values[3] >= 0.99:
        return None
    return dict(zip(('left', 'right', 'top', 'bottom'),
                    (round(value, 6) for value in values))) if any(values) else {}


def _valid_crop(value):
    if not isinstance(value, dict) or set(value) != {'left', 'right', 'top', 'bottom'}:
        return False
    if not all(type(value[k]) in (int, float) and math.isfinite(value[k])
               and 0 <= value[k] <= 0.95 for k in value):
        return False
    return value['left'] + value['right'] < 0.99 and value['top'] + value['bottom'] < 0.99


def pptx_to_portable(source, title=None, media_dir=None):
    from pptx import Presentation
    validate_pptx_archive(source)
    if title is not None and not _valid_text(title, 160):
        raise ValueError('Invalid presentation title')
    presentation = Presentation(str(source))
    activities = []
    imported_assets = []
    total_media_bytes = 0
    media_by_digest = {}  # Share identical embedded pictures across slides.
    if media_dir is not None:
        if any(parent.is_symlink() for parent in (Path(media_dir), *Path(media_dir).parents)):
            raise ValueError('Symlinked import media directory')
        Path(media_dir).mkdir(parents=True, exist_ok=True, mode=0o700)
    if not 1 <= len(presentation.slides) <= 200:
        raise ValueError('Quorivio supports at most 200 pages per import')
    for index, slide in enumerate(presentation.slides):
        widgets = []
        slide_heading = None
        for shape in _flatten_shapes(slide.shapes):
            if shape.shape_type == 13 and media_dir is not None and len(widgets) < 16:
                image = shape.image
                mime = image.content_type
                blob = image.blob
                # Validate MIME and bytes before applying deduplication or media budgets.
                # The previous mixed `and/or` condition could accept unsupported
                # formats whenever an unused media budget was available.
                supported = _supported_image(blob, mime)
                digest = hashlib.sha256(blob).hexdigest() if supported else None
                existing = media_by_digest.get((digest, mime)) if supported else None
                within_budget = (len(imported_assets) < 200 and
                                 total_media_bytes + len(blob) <= 50 * 1024 * 1024)
                if supported and (existing is not None or within_budget):
                    existing = media_by_digest.get((digest, mime))
                    if existing is None:
                        asset_id = str(uuid.uuid4())
                        with (Path(media_dir) / asset_id).open('xb') as output:
                            os.chmod(output.name, 0o600)
                            output.write(blob)
                        total_media_bytes += len(blob)
                        imported_assets.append({'id': asset_id, 'mime': mime,
                                                'bytes': len(blob), 'sha256': digest})
                        media_by_digest[(digest, mime)] = asset_id
                    else:
                        asset_id = existing
                    position = _position(shape, presentation.slide_width, presentation.slide_height)
                    rotation = float(shape.rotation or 0)
                    if math.isfinite(rotation) and rotation:
                        position['rotation'] = round(rotation % 360, 2)
                    crop = _picture_crop(shape)
                    widget = {'id': f'pptx_{index+1}_{len(widgets)+1}', 'kind': 'image',
                              'label': f'PowerPoint image {len(widgets)+1}',
                              'assetId': asset_id, 'position': position}
                    if crop:
                        widget['crop'] = crop
                    widgets.append(widget)
                continue
            # Flatten native PPTX tables to independently editable cell text widgets.
            # Preserve cell coordinates rather than discarding textual information.
            # Table semantics/borders remain explicitly lossy in fidelity reporting.
            if shape.has_table:
                table = shape.table
                left, top = int(shape.left), int(shape.top)
                row_y = top
                for row in table.rows:
                    col_x = left
                    for column_index, cell in enumerate(row.cells):
                        cell_width = int(table.columns[column_index].width)
                        content = cell.text.strip()
                        if content and len(widgets) < 16:
                            class _CellGeometry:
                                pass
                            geometry = _CellGeometry()
                            geometry.left, geometry.top = col_x, row_y
                            geometry.width, geometry.height = cell_width, int(row.height)
                            position = _position(geometry, presentation.slide_width, presentation.slide_height)
                            widgets.append({'id': f'pptx_{index+1}_{len(widgets)+1}',
                                            'kind': 'text', 'label': content[:240],
                                            'description': content[:500], 'position': position,
                                            'style': {}})
                            if slide_heading is None:
                                slide_heading = content.splitlines()[0][:160]
                        col_x += cell_width
                    row_y += int(row.height)
                continue
            # Preserve chart titles as editable text, without claiming to preserve
            # chart series, axes or styling. Charts without a title remain reported.
            if shape.has_chart:
                chart = shape.chart
                chart_title = ''
                if chart.has_title:
                    chart_title = chart.chart_title.text_frame.text.strip()
                if chart_title and len(widgets) < 16:
                    position = _position(shape, presentation.slide_width, presentation.slide_height)
                    widgets.append({'id': f'pptx_{index+1}_{len(widgets)+1}',
                                    'kind': 'text', 'label': chart_title[:240],
                                    'description': chart_title[:500],
                                    'position': position, 'style': {}})
                    if slide_heading is None:
                        slide_heading = chart_title.splitlines()[0][:160]
                continue
            from pptx.enum.shapes import MSO_SHAPE_TYPE, MSO_AUTO_SHAPE_TYPE
            editable_shape = (shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and
                              shape.auto_shape_type in (MSO_AUTO_SHAPE_TYPE.RECTANGLE,
                                  MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, MSO_AUTO_SHAPE_TYPE.OVAL,
                                  MSO_AUTO_SHAPE_TYPE.ISOSCELES_TRIANGLE))
            if not editable_shape and (not shape.has_text_frame or not shape.text.strip()):
                continue
            # Explicitly limit to supported editable elements. Warn separately.
            if len(widgets) >= 16:
                break
            # Retain author-intended whitespace and line breaks; trim only headings.
            text = shape.text if shape.has_text_frame else ''
            if slide_heading is None:
                slide_heading = next((line.strip() for line in text.splitlines() if line.strip()), None)
            first_run = next((run for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.text.strip()), None) if shape.has_text_frame else None
            style = _text_frame_style(shape)
            if first_run is not None:
                font = first_run.font
                if font.name and font.name in SUPPORTED_FONTS: style['fontFamily'] = font.name
                if font.size: style['fontSize'] = max(10, min(96, round(font.size.pt)))
                if font.bold is not None: style['bold'] = bool(font.bold)
                if font.italic is not None: style['italic'] = bool(font.italic)
                if font.underline is not None: style['underline'] = bool(font.underline)
                try:
                    if font.color and font.color.type is not None and font.color.rgb is not None:
                        style['color'] = '#' + str(font.color.rgb)
                except (AttributeError, TypeError, ValueError):
                    pass  # Theme colors cannot be represented as a fixed RGB value.
            position = _position(shape, presentation.slide_width, presentation.slide_height)
            rotation = float(shape.rotation or 0)
            if math.isfinite(rotation) and rotation:
                position['rotation'] = round(rotation % 360, 2)
            # Supported auto-shapes, including decorative shapes without text,
            # become independent editable canvas widgets.
            from pptx.enum.shapes import MSO_SHAPE_TYPE
            from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
            widget_kind = 'postit' if editable_shape else 'text'
            if widget_kind == 'postit':
                style.update(_native_shape_paint(shape))
                if shape.auto_shape_type == MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE:
                    style['rounded'] = True
                elif shape.auto_shape_type == MSO_AUTO_SHAPE_TYPE.OVAL:
                    style['shapeType'] = 'ellipse'
                elif shape.auto_shape_type == MSO_AUTO_SHAPE_TYPE.ISOSCELES_TRIANGLE:
                    style['shapeType'] = 'triangle'
            widgets.append({'id': f'pptx_{index+1}_{len(widgets)+1}', 'kind': widget_kind,
                            'label': text[:240], 'description': text[:500],
                            'position': position, 'style': style})
        page_design = {'widgets': widgets}
        background_color = _slide_background(slide)
        if background_color is not None:
            page_design['backgroundColor'] = background_color
        activities.append({'type': 'slide', 'title': (slide_heading or f'Slide {index+1}')[:160], 'options': [],
                           'correct_option_id': None, 'correct_option_ids': [],
                           'answer_mode': 'text', 'answer_required': False,
                           'duration_seconds': None, 'page_design': page_design,
                           'result_design': {}})
    return {'import_assets': imported_assets, 'format': 'quorivio-session', 'version': 1,
            'session': {'title': (title or Path(source).stem)[:160],
                        'reference_date': date.today().isoformat(), 'theme': {},
                        'allow_anonymous': True, 'show_results': False},
            'activities': activities}


def inspect_pptx_fidelity(source):
    """Report exactly which imported PowerPoint elements cannot become editable widgets.

    This intentionally returns metadata only: never media bytes or slide text.
    """
    from pptx import Presentation
    validate_pptx_archive(source)
    deck = Presentation(str(source))
    if not 1 <= len(deck.slides) <= 200:
        raise ValueError('PowerPoint page count out of range')
    report = {'slide_count': len(deck.slides), 'editable_text_boxes': 0, 'editable_decorative_shapes': 0,
              'unsupported_pictures': 0, 'editable_pictures': 0, 'unsupported_tables': 0,
              'unsupported_charts': 0, 'editable_chart_titles': 0, 'chart_titles_over_page_limit': 0, 'unsupported_other_shapes': 0,
              'editable_table_cells': 0, 'table_cells_over_page_limit': 0,
              'text_boxes_over_page_limit': 0, 'truncated_text_boxes': 0, 'non_16_9_aspect_ratio': False,
              'off_canvas_shapes': 0, 'flattened_groups': 0,
              'editable_solid_backgrounds': 0, 'unsupported_background_fills': 0,
              'editable_text_alignments': 0, 'editable_vertical_alignments': 0,
              'mixed_paragraph_alignments': 0, 'mixed_run_styles': 0,
              'unsupported_theme_text_colors': 0, 'unsupported_shape_fills': 0,
              'unsupported_shape_outlines': 0, 'cropped_pictures': 0,
              'editable_picture_crops': 0, 'unsupported_picture_crops': 0,
              'text_with_unsupported_length': 0}
    report['non_16_9_aspect_ratio'] = abs(deck.slide_width / deck.slide_height - 16 / 9) > 0.01
    media_seen = set()
    media_total = 0
    for slide in deck.slides:
        from pptx.enum.dml import MSO_FILL
        fill_type = slide.background.fill.type
        if _slide_background(slide):
            report['editable_solid_backgrounds'] += 1
        elif fill_type not in (None, MSO_FILL.BACKGROUND):
            report['unsupported_background_fills'] += 1
        # A group is rendered as editable independent elements, not a native
        # grouping relationship. Explicitly account for that lossy conversion.
        def group_count(shapes, depth=0):
            if depth > 12:
                raise ValueError('PowerPoint groups exceed maximum nesting depth')
            from pptx.enum.shapes import MSO_SHAPE_TYPE
            return sum(1 + group_count(item.shapes, depth + 1) for item in shapes
                       if item.shape_type == MSO_SHAPE_TYPE.GROUP)
        report['flattened_groups'] += group_count(slide.shapes)
        retained = 0
        for shape in _flatten_shapes(slide.shapes):
            if (shape.left < 0 or shape.top < 0 or
                    shape.left + shape.width > deck.slide_width or
                    shape.top + shape.height > deck.slide_height):
                report['off_canvas_shapes'] += 1
            from pptx.enum.shapes import MSO_SHAPE_TYPE, MSO_AUTO_SHAPE_TYPE
            supported_shape = (shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and
                               shape.auto_shape_type in (MSO_AUTO_SHAPE_TYPE.RECTANGLE,
                                   MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
                                   MSO_AUTO_SHAPE_TYPE.OVAL,
                                   MSO_AUTO_SHAPE_TYPE.ISOSCELES_TRIANGLE))
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    if any(abs(float(getattr(shape, side, 0) or 0)) > 1e-6 for side in ('crop_left','crop_right','crop_top','crop_bottom')):
                        report['cropped_pictures'] += 1
                    image = shape.image
                    blob, mime = image.blob, image.content_type
                    supported = _supported_image(blob, mime)
                    digest = (hashlib.sha256(blob).hexdigest(), mime) if supported else None
                    within_budget = digest in media_seen or (len(media_seen) < 200 and
                                     media_total + len(blob) <= 50 * 1024 * 1024)
                except (AttributeError, TypeError, ValueError):
                    supported, within_budget = False, False
                if supported and within_budget and retained < 16:
                    if _picture_crop(shape) is None:
                        report['unsupported_picture_crops'] += 1
                    elif _picture_crop(shape):
                        report['editable_picture_crops'] += 1
                    report['editable_pictures'] += 1
                    retained += 1
                    if digest not in media_seen:
                        media_seen.add(digest)
                        media_total += len(blob)
                else:
                    report['unsupported_pictures'] += 1
            elif supported_shape or (shape.has_text_frame and shape.text.strip()):
                frame_style = _text_frame_style(shape)
                report['editable_text_alignments'] += int('textAlign' in frame_style and retained < 16)
                report['editable_vertical_alignments'] += int('verticalAlign' in frame_style and retained < 16)
                if shape.has_text_frame and len({p.alignment for p in shape.text_frame.paragraphs if p.text.strip()}) > 1:
                    report['mixed_paragraph_alignments'] += 1
                if shape.has_text_frame and len(shape.text) > 500:
                    report['truncated_text_boxes'] += 1
                    report['text_with_unsupported_length'] += 1
                if shape.has_text_frame:
                    runs = [r for para in shape.text_frame.paragraphs for r in para.runs if r.text]
                    if runs:
                        def _run_style(run):
                            font = run.font
                            try:
                                color = (str(font.color.type), str(font.color.rgb))
                            except (AttributeError, TypeError, ValueError):
                                color = (str(font.color.type), None)
                            return (font.name, font.size.pt if font.size else None,
                                    font.bold, font.italic, font.underline, color)
                        if len({_run_style(run) for run in runs}) > 1:
                            report['mixed_run_styles'] += 1
                        from pptx.enum.dml import MSO_COLOR_TYPE
                        for run in runs:
                            if run.font.color.type in (MSO_COLOR_TYPE.SCHEME, MSO_COLOR_TYPE.SYSTEM):
                                report['unsupported_theme_text_colors'] += 1
                                break
                if supported_shape:
                    try:
                        if shape.fill.type not in (None, MSO_FILL.SOLID, MSO_FILL.BACKGROUND):
                            report['unsupported_shape_fills'] += 1
                    except (AttributeError, TypeError, ValueError):
                        pass
                    try:
                        from pptx.enum.dml import MSO_COLOR_TYPE
                        from pptx.enum.dml import MSO_LINE_DASH_STYLE
                        from pptx.enum.dml import MSO_FILL
                        line = shape.line
                        unsupported = (line.fill.type not in (None, MSO_FILL.BACKGROUND) and
                                       line.color.type not in (None, MSO_COLOR_TYPE.RGB))
                        if line.color.type == MSO_COLOR_TYPE.RGB and line.fill.type != MSO_FILL.BACKGROUND:
                            unsupported |= bool(line.width) and not (0.25 <= line.width.pt <= 12)
                            unsupported |= line.dash_style not in (None, MSO_LINE_DASH_STYLE.SOLID,
                                MSO_LINE_DASH_STYLE.DASH, MSO_LINE_DASH_STYLE.ROUND_DOT)
                        if unsupported:
                            report['unsupported_shape_outlines'] += 1
                    except (AttributeError, TypeError, ValueError):
                        pass
                if retained < 16:
                    if supported_shape and not (shape.has_text_frame and shape.text.strip()):
                        report['editable_decorative_shapes'] += 1
                    else:
                        report['editable_text_boxes'] += 1
                    retained += 1
                else:
                    report['text_boxes_over_page_limit'] += 1
            elif shape.has_table:
                # The native table structure, styles and merging are not retained,
                # even when its cell text is imported as editable elements.
                report['unsupported_tables'] += 1
                for row in shape.table.rows:
                    for cell in row.cells:
                        if not cell.text.strip():
                            continue
                        if retained < 16:
                            report['editable_table_cells'] += 1
                            retained += 1
                        else:
                            report['table_cells_over_page_limit'] += 1
            elif shape.has_chart:
                # Native data series are always lossy. The title can be recovered
                # as editable text when the page still has capacity.
                report['unsupported_charts'] += 1
                try:
                    has_title = shape.chart.has_title and bool(shape.chart.chart_title.text_frame.text.strip())
                except (AttributeError, ValueError):
                    has_title = False
                if has_title:
                    if retained < 16:
                        retained += 1
                        report['editable_chart_titles'] += 1
                    else:
                        report['chart_titles_over_page_limit'] += 1
            else:
                report['unsupported_other_shapes'] += 1
    report['lossy'] = any(report[k] for k in (
        'unsupported_pictures', 'unsupported_tables', 'unsupported_charts',
        'unsupported_other_shapes', 'text_boxes_over_page_limit', 'table_cells_over_page_limit', 'flattened_groups',
        'non_16_9_aspect_ratio', 'off_canvas_shapes', 'truncated_text_boxes',
        'unsupported_background_fills', 'mixed_paragraph_alignments',
        'mixed_run_styles', 'unsupported_theme_text_colors', 'unsupported_shape_fills',
        'unsupported_shape_outlines', 'unsupported_picture_crops'))
    return report


def portable_to_pptx(data, destination, media_dir=None):
    from pptx import Presentation
    from pptx.util import Inches
    activities = _validate_export_input(data)
    p = Presentation()
    p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
    seen_assets = set()
    total_asset_bytes = 0
    for activity in activities:
        slide = p.slides.add_slide(p.slide_layouts[6])
        design = activity.get('page_design') or {}
        background = design.get('backgroundColor')
        if background is not None:
            from pptx.dml.color import RGBColor
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor.from_string(background[1:])
        widgets = (activity.get('page_design') or {}).get('widgets') or []
        if len(widgets) > 16:
            raise ValueError('Too many widgets')
        for widget in widgets:
            if widget.get('kind') not in ('text', 'postit', 'question', 'timed-quiz', 'image', 'logo'):
                continue
            position = widget.get('position') or {'x': 5, 'y': 5, 'width': 90, 'height': 20}
            x, y, w, h = (float(position.get(k, v)) for k, v in [('x', 5), ('y', 5), ('width', 90), ('height', 20)])
            if not all(math.isfinite(v) for v in (x, y, w, h)) or x < 0 or y < 0 or w <= 0 or h <= 0 or x+w > 100.00001 or y+h > 100.00001:
                raise ValueError('Invalid widget geometry')
            if widget.get('kind') in ('image', 'logo'):
                asset_id = widget.get('assetId')
                try:
                    valid_asset_id = isinstance(asset_id, str) and str(uuid.UUID(asset_id)) == asset_id and uuid.UUID(asset_id).version == 4
                except (ValueError, AttributeError, TypeError):
                    valid_asset_id = False
                if not valid_asset_id:
                    raise ValueError('Invalid image asset ID')
                if media_dir is None:
                    raise ValueError('Image export requires a private media directory')
                picture_path = Path(media_dir) / asset_id
                if (any(parent.is_symlink() for parent in (Path(media_dir), *Path(media_dir).parents)) or
                        picture_path.is_symlink() or
                        not picture_path.is_file() or picture_path.stat().st_size > 10 * 1024 * 1024):
                    raise ValueError('Missing or oversized image asset')
                blob = picture_path.read_bytes()
                if not (_supported_image(blob, 'image/png') or _supported_image(blob, 'image/jpeg')):
                    raise ValueError('Unsupported or invalid image asset')
                if asset_id not in seen_assets:
                    seen_assets.add(asset_id)
                    total_asset_bytes += len(blob)
                    if len(seen_assets) > 200 or total_asset_bytes > 50 * 1024 * 1024:
                        raise ValueError('PowerPoint media quota exceeded')
                from pptx.enum.shapes import MSO_SHAPE_TYPE
                # Serialize exactly the verified bytes: never reopen a mutable asset path.
                shape = slide.shapes.add_picture(io.BytesIO(blob), int(p.slide_width*x/100),
                                                 int(p.slide_height*y/100), int(p.slide_width*w/100),
                                                 int(p.slide_height*h/100))
                crop = widget.get('crop')
                if crop is not None:
                    if not _valid_crop(crop):
                        raise ValueError('Invalid picture crop')
                    for side, fraction in crop.items():
                        setattr(shape, 'crop_' + side, fraction)
                rotation = position.get('rotation', 0)
                if isinstance(rotation, (int, float)) and math.isfinite(rotation):
                    shape.rotation = rotation % 360
                continue
            if widget.get('kind') == 'postit':
                from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
                style = widget.get('style') or {}
                box = slide.shapes.add_shape(
                    ({'ellipse': MSO_AUTO_SHAPE_TYPE.OVAL, 'triangle': MSO_AUTO_SHAPE_TYPE.ISOSCELES_TRIANGLE}.get(style.get('shapeType'))
                     or (MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if style.get('rounded') is True else MSO_AUTO_SHAPE_TYPE.RECTANGLE)),
                    int(p.slide_width*x/100), int(p.slide_height*y/100),
                    int(p.slide_width*w/100), int(p.slide_height*h/100))
            else:
                box = slide.shapes.add_textbox(int(p.slide_width*x/100), int(p.slide_height*y/100),
                    int(p.slide_width*w/100), int(p.slide_height*h/100))
            content = str(widget.get('label') or '')
            # Imports keep the short UI label and a longer editable description.
            # Export the full text once, instead of repeating the truncated label.
            description = widget.get('description')
            if widget.get('kind') in ('text', 'postit') and isinstance(description, str) and description.startswith(content):
                content = description
            if widget.get('kind') in ('question', 'timed-quiz'):
                content += '\n' + '\n'.join(str(o.get('label', '')) for o in widget.get('options', [])[:12])
            elif widget.get('description') and widget.get('description') != content and widget.get('kind') not in ('text', 'postit'):
                content += '\n' + str(widget['description'])
            box.text = content[:MAX_EXPORT_TEXT]
            if widget.get('kind') == 'postit' and isinstance(widget.get('style'), dict):
                from pptx.dml.color import RGBColor
                colors = widget['style']
                if colors.get('fillNone') is True:
                    box.fill.background()
                elif colors.get('fillColor'):
                    box.fill.solid()
                    box.fill.fore_color.rgb = RGBColor.from_string(colors['fillColor'][1:])
                if colors.get('strokeNone') is True:
                    box.line.fill.background()
                elif colors.get('strokeColor'):
                    from pptx.enum.dml import MSO_LINE_DASH_STYLE
                    from pptx.util import Pt
                    box.line.color.rgb = RGBColor.from_string(colors['strokeColor'][1:])
                    if colors.get('strokeWidth') is not None:
                        box.line.width = Pt(colors['strokeWidth'])
                    if colors.get('strokeDash'):
                        box.line.dash_style = {'solid': MSO_LINE_DASH_STYLE.SOLID,
                            'dash': MSO_LINE_DASH_STYLE.DASH,
                            'dot': MSO_LINE_DASH_STYLE.ROUND_DOT}[colors['strokeDash']]
            rotation = position.get('rotation', 0)
            if isinstance(rotation, (int, float)) and math.isfinite(rotation):
                box.rotation = rotation % 360
            style = widget.get('style') or {}
            if isinstance(style, dict):
                from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
                if style.get('textAlign'):
                    alignment = {'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[style['textAlign']]
                    for paragraph in box.text_frame.paragraphs:
                        paragraph.alignment = alignment
                if style.get('verticalAlign'):
                    box.text_frame.vertical_anchor = {'top':MSO_ANCHOR.TOP,'middle':MSO_ANCHOR.MIDDLE,'bottom':MSO_ANCHOR.BOTTOM}[style['verticalAlign']]
                from pptx.util import Pt
                from pptx.dml.color import RGBColor
                for paragraph in box.text_frame.paragraphs:
                    for run in paragraph.runs:
                        font = run.font
                        if isinstance(style.get('fontFamily'), str) and style['fontFamily']:
                            font.name = style['fontFamily'][:80]
                        size = style.get('fontSize')
                        if isinstance(size, (int, float)) and math.isfinite(size):
                            font.size = Pt(max(10, min(96, size)))
                        if isinstance(style.get('bold'), bool): font.bold = style['bold']
                        if isinstance(style.get('italic'), bool): font.italic = style['italic']
                        if isinstance(style.get('underline'), bool): font.underline = style['underline']
                        color = style.get('color')
                        if isinstance(color, str) and len(color) == 7 and color.startswith('#'):
                            try: font.color.rgb = RGBColor.from_string(color[1:])
                            except ValueError: pass
    destination = Path(destination)
    if destination.exists() or destination.is_symlink():
        raise FileExistsError('PowerPoint destination already exists')
    if any(parent.is_symlink() for parent in destination.parents):
        raise ValueError('Symlinked PowerPoint destination parent')
    stage = None
    try:
        with tempfile.NamedTemporaryFile(prefix='.quorivio-pptx-', suffix='.tmp',
                                         dir=destination.parent, delete=False) as stream:
            stage = Path(stream.name)
            os.chmod(stage, 0o600)
        p.save(str(stage))
        if stage.stat().st_size > MAX_PPTX_BYTES:
            raise ValueError('Exported PowerPoint exceeds size limit')
        # Hard-link creation is atomic and fails closed on a competing publisher.
        os.link(stage, destination, follow_symlinks=False)
        stage.unlink()
        stage = None
    finally:
        if stage is not None:
            stage.unlink(missing_ok=True)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest='operation', required=True)
    for command in ('import', 'export'):
        action = sub.add_parser(command)
        action.add_argument('source')
        action.add_argument('destination')
    args = parser.parse_args()
    if args.operation == 'import':
        destination = Path(args.destination)
        if destination.exists() or destination.is_symlink():
            raise FileExistsError('Portable destination already exists')
        if any(parent.is_symlink() for parent in destination.parents):
            raise ValueError('Symlinked portable destination parent')
        payload = json.dumps(pptx_to_portable(args.source), ensure_ascii=False, indent=2, allow_nan=False).encode('utf-8')
        if len(payload) > 2 * 1024 * 1024:
            raise ValueError('Portable definition too large')
        stage = None
        try:
            with tempfile.NamedTemporaryFile(prefix='.quorivio-portable-', suffix='.tmp',
                                             dir=destination.parent, delete=False) as stream:
                stage = Path(stream.name)
                os.chmod(stage, 0o600)
                stream.write(payload)
            os.link(stage, destination, follow_symlinks=False)
            stage.unlink()
            stage = None
        finally:
            if stage is not None:
                stage.unlink(missing_ok=True)
    else:
        portable_to_pptx(json.loads(Path(args.source).read_text(encoding='utf-8')), args.destination)


if __name__ == '__main__':
    main()
