"""30 independent acceptance cases: editable native PPTX picture crop + visual roundtrip."""
import copy
import io
import math
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from pptx_portable import _picture_crop, _valid_crop, inspect_pptx_fidelity, portable_to_pptx, pptx_to_portable
from pptx import Presentation
from pptx.util import Inches
from PIL import Image, ImageChops, ImageStat, ImageDraw

CROPS = [
    (0.10, 0, 0, 0), (0, 0.15, 0, 0), (0, 0, 0.2, 0), (0, 0, 0, 0.25),
    (0.1, 0.2, 0, 0), (0, 0, 0.15, 0.2), (0.15, 0, 0.1, 0),
    (0, 0.12, 0, 0.18), (0.1, 0.2, 0.15, 0.1),
    (0.3, 0.15, 0.2, 0.1), (0.05, 0.35, 0.3, 0.05), (0.22, 0.18, 0.12, 0.16),
]


class SeventhConversion30(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory(prefix='quorivio-picture-crop-')
        cls.root = Path(cls.tmp.name)
        img = Image.new('RGB', (480, 320), '#f8e4d0')
        painter = ImageDraw.Draw(img)
        for x in range(0, 480, 40):
            for y in range(0, 320, 40):
                painter.rectangle((x, y, x+39, y+39), fill=(30+x//3, 40+y//2, 90+(x+y)//8))
        painter.ellipse((100, 60, 380, 260), outline='white', width=10)
        cls.png = cls.root/'pattern.png'
        img.save(cls.png)
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
        for i, fractions in enumerate(CROPS):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            shape = slide.shapes.add_picture(str(cls.png), Inches(1.25), Inches(0.9), Inches(10.5), Inches(5.6))
            for side, value in zip(('left', 'right', 'top', 'bottom'), fractions):
                setattr(shape, 'crop_'+side, value)
        cls.source = cls.root/'source.pptx'
        cls.export = cls.root/'export.pptx'
        deck.save(cls.source)
        cls.media = cls.root/'media'
        cls.portable = pptx_to_portable(cls.source, media_dir=cls.media)
        portable_to_pptx(cls.portable, cls.export, media_dir=cls.media)
        cls.report = inspect_pptx_fidelity(cls.source)
        cls.render_available = all(shutil.which(name) for name in ('libreoffice', 'pdftoppm'))
        cls.rendered = {}
        if cls.render_available:
            for pptx in (cls.source, cls.export):
                profile = cls.root/('lo-'+pptx.stem)
                result = subprocess.run(['libreoffice', '-env:UserInstallation=file://'+str(profile),
                    '--headless', '--convert-to', 'pdf', '--outdir', str(cls.root), str(pptx)],
                    capture_output=True, text=True, timeout=90)
                if result.returncode or not pptx.with_suffix('.pdf').is_file():
                    raise AssertionError('LibreOffice failed: '+result.stderr)
                subprocess.run(['pdftoppm', '-f', '1', '-l', '12', '-scale-to', '480', '-png',
                    str(pptx.with_suffix('.pdf')), str(cls.root/pptx.stem)],
                    check=True, capture_output=True, timeout=90)
                cls.rendered[pptx.stem] = sorted(cls.root.glob(pptx.stem+'-*.png'))

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def _visual(self, index):
        self.assertTrue(self.render_available, 'LibreOffice/Poppler are required')
        self.assertEqual(len(self.rendered['source']), 12)
        self.assertEqual(len(self.rendered['export']), 12)
        with Image.open(self.rendered['source'][index]) as original, Image.open(self.rendered['export'][index]) as exported:
            self.assertEqual(original.size, exported.size)
            difference = ImageChops.difference(original.convert('RGB'), exported.convert('RGB'))
            self.assertLess(sum(ImageStat.Stat(difference).mean)/3, 14, f'Crop visual drift on slide {index+1}')

    def test_01_visual_crop_left(self): self._visual(0)
    def test_02_visual_crop_right(self): self._visual(1)
    def test_03_visual_crop_top(self): self._visual(2)
    def test_04_visual_crop_bottom(self): self._visual(3)
    def test_05_visual_crop_horizontal(self): self._visual(4)
    def test_06_visual_crop_vertical(self): self._visual(5)
    def test_07_visual_crop_left_top(self): self._visual(6)
    def test_08_visual_crop_right_bottom(self): self._visual(7)
    def test_09_visual_crop_all_sides(self): self._visual(8)
    def test_10_visual_crop_large(self): self._visual(9)
    def test_11_visual_crop_asymmetric(self): self._visual(10)
    def test_12_visual_crop_centered(self): self._visual(11)

    def test_13_all_twelve_slides_keep_editable_image(self):
        self.assertTrue(all(len(a['page_design']['widgets']) == 1 and
            a['page_design']['widgets'][0]['kind'] == 'image' for a in self.portable['activities']))

    def test_14_all_crop_values_imported(self):
        for i, values in enumerate(CROPS):
            crop = self.portable['activities'][i]['page_design']['widgets'][0]['crop']
            for side, expected in zip(('left', 'right', 'top', 'bottom'), values):
                self.assertAlmostEqual(crop[side], expected, places=5)

    def test_15_all_crop_values_exported(self):
        for i, slide in enumerate(Presentation(self.export).slides):
            for side, expected in zip(('left', 'right', 'top', 'bottom'), CROPS[i]):
                self.assertAlmostEqual(getattr(slide.shapes[0], 'crop_'+side), expected, places=5)

    def test_16_media_is_deduplicated_across_crop_variants(self):
        self.assertEqual(len(self.portable['import_assets']), 1)
        self.assertEqual(len({a['page_design']['widgets'][0]['assetId'] for a in self.portable['activities']}), 1)

    def test_17_cropped_pictures_no_longer_count_as_lossy(self):
        self.assertEqual(self.report['cropped_pictures'], 12)
        self.assertEqual(self.report['editable_picture_crops'], 12)
        self.assertEqual(self.report['unsupported_picture_crops'], 0)
        self.assertFalse(self.report['lossy'], self.report)

    def test_18_exported_pictures_are_native_shapes(self):
        self.assertTrue(all(slide.shapes[0].shape_type == 13 for slide in Presentation(self.export).slides))

    def test_19_zero_crop_is_valid(self):
        self.assertTrue(_valid_crop(dict.fromkeys(('left','right','top','bottom'), 0)))

    def test_20_legal_crop_upper_boundary(self):
        self.assertTrue(_valid_crop({'left':.95,'right':0,'top':0,'bottom':.95}))

    def test_21_reject_horizontal_overcrop(self):
        self.assertFalse(_valid_crop({'left':.5,'right':.5,'top':0,'bottom':0}))

    def test_22_reject_vertical_overcrop(self):
        self.assertFalse(_valid_crop({'left':0,'right':0,'top':.7,'bottom':.3}))

    def test_23_reject_negative_crop(self):
        self.assertFalse(_valid_crop({'left':-.01,'right':0,'top':0,'bottom':0}))

    def test_24_reject_nonfinite_crop(self):
        for bad in (float('nan'),float('inf'),float('-inf')):
            self.assertFalse(_valid_crop({'left':bad,'right':0,'top':0,'bottom':0}))

    def test_25_reject_boolean_crop(self):
        self.assertFalse(_valid_crop({'left':True,'right':0,'top':0,'bottom':0}))

    def test_26_reject_extra_crop_fields(self):
        self.assertFalse(_valid_crop({'left':0,'right':0,'top':0,'bottom':0,'css':'url()'}))

    def test_27_reject_missing_crop_fields(self):
        self.assertFalse(_valid_crop({'left':0,'right':0,'top':0}))

    def test_28_reject_export_invalid_crop_before_file_creation(self):
        bad = copy.deepcopy(self.portable)
        bad['activities'][0]['page_design']['widgets'][0]['crop']['left'] = 1
        destination = self.root/'rejected.pptx'
        with self.assertRaisesRegex(ValueError, 'Invalid picture crop'):
            portable_to_pptx(bad, destination, media_dir=self.media)
        self.assertFalse(destination.exists())

    def test_29_reject_crop_on_text_widget(self):
        bad = copy.deepcopy(self.portable)
        widget = bad['activities'][0]['page_design']['widgets'][0]
        widget['kind'] = 'text'
        with self.assertRaisesRegex(ValueError, 'Crop is only supported'):
            portable_to_pptx(bad, self.root/'rejected-text.pptx')

    def test_30_unsupported_negative_native_crop_is_reported(self):
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        shape = slide.shapes.add_picture(str(self.png), Inches(1), Inches(1), Inches(4), Inches(3))
        shape.crop_left = -.1
        source = self.root/'overscan.pptx'
        deck.save(source)
        report = inspect_pptx_fidelity(source)
        self.assertEqual(report['unsupported_picture_crops'], 1)
        self.assertTrue(report['lossy'])
        self.assertIsNone(_picture_crop(Presentation(source).slides[0].shapes[0]))


if __name__ == '__main__':
    unittest.main()
