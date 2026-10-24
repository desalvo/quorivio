"""Native PowerPoint table text is recoverable as individually editable widgets."""
import tempfile
import unittest
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from tools.pptx_portable import pptx_to_portable, inspect_pptx_fidelity

class PptxEditableTableCellsTest(unittest.TestCase):
    def build(self, rows, cols):
        temp = tempfile.TemporaryDirectory()
        path = Path(temp.name) / 'table.pptx'
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        table = slide.shapes.add_table(rows, cols, Inches(1), Inches(1), Inches(8), Inches(3)).table
        for r in range(rows):
            for c in range(cols):
                table.cell(r,c).text = f'Cell {r}-{c}'
        deck.save(path)
        return temp,path

    def test_table_cells_are_editable_but_fidelity_reports_loss(self):
        tmp,path = self.build(2,3)
        try:
            definition=pptx_to_portable(path)
            widgets=definition['activities'][0]['page_design']['widgets']
            self.assertEqual(len(widgets),6)
            self.assertTrue(all(w['kind']=='text' for w in widgets))
            self.assertEqual(widgets[0]['label'],'Cell 0-0')
            self.assertGreater(widgets[1]['position']['x'],widgets[0]['position']['x'])
            self.assertGreater(widgets[3]['position']['y'],widgets[0]['position']['y'])
            report=inspect_pptx_fidelity(path)
            self.assertEqual(report['editable_table_cells'],6)
            self.assertEqual(report['unsupported_tables'],1)
            self.assertTrue(report['lossy'])
        finally: tmp.cleanup()

    def test_table_cell_limit_matches_import_and_report(self):
        tmp,path=self.build(5,4)
        try:
            widgets=pptx_to_portable(path)['activities'][0]['page_design']['widgets']
            report=inspect_pptx_fidelity(path)
            self.assertEqual(len(widgets),16)
            self.assertEqual(report['editable_table_cells'],16)
            self.assertEqual(report['table_cells_over_page_limit'],4)
            self.assertTrue(report['lossy'])
        finally: tmp.cleanup()

if __name__=='__main__': unittest.main()
