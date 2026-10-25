"""30 independently asserted conversion-fidelity acceptance cases.

Twelve pages are rendered by real LibreOffice and Poppler, and checked against
original renders; remaining tests cover editable OOXML/portable fidelity and
honest reporting of unsupported native PowerPoint constructs.
"""
import copy
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from pptx_portable import inspect_pptx_fidelity, portable_to_pptx, pptx_to_portable
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_COLOR_TYPE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from PIL import Image, ImageChops, ImageStat

COLORS = ['F5E5D5','D5E5F5','E5F5D5','F5D5E5','D5F5E5','E5D5F5',
          'E5E5E5','F5F5D5','D5F5F5','F5D5D5','D5D5F5','F5E5E5']


class ConversionSixth30(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory(prefix='quorivio-fidelity-')
        cls.root = Path(cls.tmp.name)
        cls.deck = Presentation()
        cls.deck.slide_width = Inches(13.333)
        cls.deck.slide_height = Inches(7.5)
        for i, color in enumerate(COLORS):
            slide = cls.deck.slides.add_slide(cls.deck.slide_layouts[6])
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor.from_string(color)
            shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(9), Inches(1.5))
            shape.text = f'Visual slide {i + 1:02d}'
            shape.text_frame.paragraphs[0].alignment = (PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.RIGHT)[i % 3]
            run = shape.text_frame.paragraphs[0].runs[0]
            run.font.name = 'Arial'
            run.font.size = Pt(32)
            run.font.bold = i % 2 == 0
            run.font.color.rgb = RGBColor(20, 20, 20)
            rectangle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE,
                        Inches(1.5), Inches(3), Inches(2), Inches(1.3))
            rectangle.fill.solid()
            rectangle.fill.fore_color.rgb = RGBColor(40, 80, 130)
            rectangle.line.color.rgb = RGBColor(20, 20, 20)
        cls.source = cls.root / 'source.pptx'
        cls.roundtrip = cls.root / 'roundtrip.pptx'
        cls.deck.save(cls.source)
        cls.portable = pptx_to_portable(cls.source)
        portable_to_pptx(cls.portable, cls.roundtrip)
        cls.original_pdf = cls.root / 'source.pdf'
        cls.roundtrip_pdf = cls.root / 'roundtrip.pdf'
        cls.render_available = all(shutil.which(cmd) for cmd in ('libreoffice','pdftoppm','pdfinfo'))
        if cls.render_available:
            for pptx in (cls.source, cls.roundtrip):
                profile = cls.root / ('lo-'+pptx.stem)
                result = subprocess.run(['libreoffice','-env:UserInstallation=file://'+str(profile),
                    '--headless','--convert-to','pdf','--outdir',str(cls.root),str(pptx)],
                    capture_output=True,text=True,timeout=90)
                if result.returncode or not pptx.with_suffix('.pdf').is_file():
                    raise AssertionError('LibreOffice PDF rendering failed: '+result.stderr)
            cls.rendered = {}
            for pdf in (cls.original_pdf, cls.roundtrip_pdf):
                prefix = cls.root / pdf.stem
                subprocess.run(['pdftoppm','-f','1','-l','12','-scale-to','480','-png',
                    str(pdf),str(prefix)],check=True,capture_output=True,timeout=90)
                cls.rendered[pdf.stem] = sorted(cls.root.glob(pdf.stem+'-*.png'))

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def _slide(self, index):
        self.assertTrue(self.render_available, 'LibreOffice and Poppler required for real visual corpus')
        originals = self.rendered['source']
        exports = self.rendered['roundtrip']
        self.assertEqual(len(originals),12)
        self.assertEqual(len(exports),12)
        with Image.open(originals[index]) as a, Image.open(exports[index]) as b:
            self.assertEqual(a.size,b.size)
            a,b=a.convert('RGB'),b.convert('RGB')
            # Background and major geometry must remain visually recognizable;
            # rasterizers can differ on font substitution and antialiasing.
            diff=ImageChops.difference(a,b)
            mean=sum(ImageStat.Stat(diff).mean)/3
            self.assertLess(mean, 18, f'Visual drift on slide {index+1}: {mean:.2f}')
            self.assertLess(max(ImageStat.Stat(b).stddev), 85)

    def _audit_fixture(self, mutator):
        deck=Presentation()
        slide=deck.slides.add_slide(deck.slide_layouts[6])
        shape=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(4),Inches(2))
        shape.text='A'
        mutator(slide,shape)
        path=self.root / 'audit.pptx'
        deck.save(path)
        return inspect_pptx_fidelity(path),pptx_to_portable(path)

    def test_01_render_slide_01(self): self._slide(0)
    def test_02_render_slide_02(self): self._slide(1)
    def test_03_render_slide_03(self): self._slide(2)
    def test_04_render_slide_04(self): self._slide(3)
    def test_05_render_slide_05(self): self._slide(4)
    def test_06_render_slide_06(self): self._slide(5)
    def test_07_render_slide_07(self): self._slide(6)
    def test_08_render_slide_08(self): self._slide(7)
    def test_09_render_slide_09(self): self._slide(8)
    def test_10_render_slide_10(self): self._slide(9)
    def test_11_render_slide_11(self): self._slide(10)
    def test_12_render_slide_12(self): self._slide(11)

    def test_13_slide_count_roundtrip(self):
        self.assertEqual(len(self.portable['activities']),12)
        self.assertEqual(len(Presentation(self.roundtrip).slides),12)

    def test_14_backgrounds_remain_editable(self):
        for i,color in enumerate(COLORS):
            self.assertEqual(self.portable['activities'][i]['page_design']['backgroundColor'], '#'+color)

    def test_15_shape_geometry_roundtrip(self):
        for slide in self.portable['activities']:
            widgets=slide['page_design']['widgets']
            self.assertEqual(len(widgets),2)
            self.assertEqual(widgets[1]['kind'],'postit')
            self.assertAlmostEqual(widgets[1]['position']['x'],11.25,delta=0.2)

    def test_16_text_alignment_roundtrip(self):
        for i,slide in enumerate(self.portable['activities']):
            self.assertEqual(slide['page_design']['widgets'][0]['style']['textAlign'],
                             ('left','center','right')[i%3])

    def test_17_bold_run_roundtrip(self):
        for i,slide in enumerate(self.portable['activities']):
            self.assertEqual(slide['page_design']['widgets'][0]['style']['bold'],i%2==0)

    def test_18_native_shape_fill_roundtrip(self):
        for slide in Presentation(self.roundtrip).slides:
            self.assertEqual(str(slide.shapes[1].fill.fore_color.rgb),'285082')

    def test_19_native_text_color_roundtrip(self):
        for slide in Presentation(self.roundtrip).slides:
            self.assertEqual(str(slide.shapes[0].text_frame.paragraphs[0].runs[0].font.color.rgb),'141414')

    def test_20_fidelity_report_supported_corpus(self):
        report=inspect_pptx_fidelity(self.source)
        self.assertEqual(report['editable_solid_backgrounds'],12)
        self.assertEqual(report['editable_text_boxes'],12)
        self.assertEqual(report['editable_decorative_shapes'],12)
        self.assertFalse(report['lossy'],report)

    def test_21_leading_whitespace_preserved(self):
        _, portable=self._audit_fixture(lambda s,x: setattr(x,'text','  Leading'))
        self.assertEqual(portable['activities'][0]['page_design']['widgets'][0]['description'],'  Leading')

    def test_22_trailing_whitespace_preserved(self):
        _, portable=self._audit_fixture(lambda s,x: setattr(x,'text','Trailing  '))
        self.assertEqual(portable['activities'][0]['page_design']['widgets'][0]['description'],'Trailing  ')

    def test_23_linebreak_preserved(self):
        _, portable=self._audit_fixture(lambda s,x: setattr(x,'text','First\nSecond'))
        self.assertEqual(portable['activities'][0]['page_design']['widgets'][0]['description'],'First\nSecond')

    def test_24_unicode_text_preserved(self):
        _, portable=self._audit_fixture(lambda s,x: setattr(x,'text','Città — élève 日本語'))
        self.assertEqual(portable['activities'][0]['page_design']['widgets'][0]['description'],'Città — élève 日本語')

    def test_25_mixed_run_styles_detected(self):
        def mutate(s,x):
            p=x.text_frame.paragraphs[0]
            p.runs[0].font.bold=True
            p.add_run().text='B'
            p.runs[1].font.bold=False
        report,_=self._audit_fixture(mutate)
        self.assertEqual(report['mixed_run_styles'],1)
        self.assertTrue(report['lossy'])

    def test_26_uniform_run_styles_not_misreported(self):
        def mutate(s,x):
            p=x.text_frame.paragraphs[0]
            p.runs[0].font.bold=True
            r=p.add_run();r.text='B';r.font.bold=True
        report,_=self._audit_fixture(mutate)
        self.assertEqual(report['mixed_run_styles'],0)

    def test_27_theme_text_color_detected(self):
        def mutate(s,x):
            from pptx.enum.dml import MSO_THEME_COLOR
            x.text_frame.paragraphs[0].runs[0].font.color.theme_color=MSO_THEME_COLOR.ACCENT_1
        report,_=self._audit_fixture(mutate)
        self.assertEqual(report['unsupported_theme_text_colors'],1)
        self.assertTrue(report['lossy'])

    def test_28_gradient_shape_fill_detected(self):
        def mutate(s,x):
            from pptx.oxml.xmlchemy import OxmlElement
            rect=s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE,Inches(1),Inches(3),Inches(2),Inches(1))
            sppr=rect._element.spPr
            for element in list(sppr):
                if element.tag.endswith(('solidFill','noFill','gradFill')):sppr.remove(element)
            sppr.append(OxmlElement('a:gradFill'))
        report,_=self._audit_fixture(mutate)
        self.assertEqual(report['unsupported_shape_fills'],1)
        self.assertTrue(report['lossy'])

    def test_29_theme_shape_outline_detected(self):
        def mutate(s,x):
            from pptx.enum.dml import MSO_THEME_COLOR
            rect=s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE,Inches(1),Inches(3),Inches(2),Inches(1))
            rect.line.color.theme_color=MSO_THEME_COLOR.ACCENT_1
        report,_=self._audit_fixture(mutate)
        self.assertEqual(report['unsupported_shape_outlines'],1)
        self.assertTrue(report['lossy'])

    def test_30_long_text_truncation_explicit(self):
        report,portable=self._audit_fixture(lambda s,x: setattr(x,'text','X'*501))
        self.assertEqual(report['text_with_unsupported_length'],1)
        self.assertTrue(report['lossy'])
        self.assertEqual(len(portable['activities'][0]['page_design']['widgets'][0]['description']),500)


if __name__ == '__main__':
    unittest.main()
