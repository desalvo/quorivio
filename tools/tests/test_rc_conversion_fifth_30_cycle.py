"""Thirty real PPTX import/export fidelity and validation checks.

Fixtures are serialized through python-pptx, parsed back and checked against
editable portable page definitions; negative cases must publish no output.
"""
import copy
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from pptx_portable import pptx_to_portable, portable_to_pptx, inspect_pptx_fidelity
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches


def definition(style=None, background=None):
    design = {'widgets': [{'id': 'text-1', 'kind': 'text', 'label': 'Hello',
                           'description': 'Hello', 'style': style or {},
                           'position': {'x': 10, 'y': 10, 'width': 50, 'height': 30}}]}
    if background is not None:
        design['backgroundColor'] = background
    return {'format': 'quorivio-session', 'version': 1, 'activities': [
        {'type': 'slide', 'page_design': design}]}


class ConversionFifth30(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)

    def source(self):
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
        shape.text = 'Hello\nWorld'
        path = self.root / 'source.pptx'
        return deck, slide, shape, path

    def test_01_black_background_import(self):
        self._import_background('000000')

    def test_02_white_background_import(self):
        self._import_background('FFFFFF')

    def test_03_brand_background_import(self):
        self._import_background('12AB34')

    def test_04_red_background_import(self):
        self._import_background('FF0000')

    def test_05_blue_background_import(self):
        self._import_background('0000FF')

    def test_06_background_export_roundtrip(self):
        self._export_background('#12AB34')

    def test_07_background_black_export_roundtrip(self):
        self._export_background('#000000')

    def test_08_background_white_export_roundtrip(self):
        self._export_background('#FFFFFF')

    def test_09_background_lowercase_export_roundtrip(self):
        self._export_background('#ff10ab')

    def test_10_background_fidelity_count(self):
        deck, slide, _, path = self.source()
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(12, 34, 56)
        deck.save(path)
        report = inspect_pptx_fidelity(path)
        self.assertEqual(report['editable_solid_backgrounds'], 1)
        self.assertEqual(report['unsupported_background_fills'], 0)

    def test_11_background_absent_not_fabricated(self):
        deck, _, _, path = self.source()
        deck.save(path)
        self.assertNotIn('backgroundColor', pptx_to_portable(path)['activities'][0]['page_design'])

    def test_12_background_unsupported_gradient_reported(self):
        deck, slide, _, path = self.source()
        from pptx.oxml.xmlchemy import OxmlElement
        # OOXML gradient fill is a legal background but is not representable by a flat RGB.
        bg_pr = slide.background._element.get_or_add_bgPr()
        for child in list(bg_pr):
            bg_pr.remove(child)
        bg_pr.append(OxmlElement('a:gradFill'))
        deck.save(path)
        report = inspect_pptx_fidelity(path)
        self.assertEqual(report['unsupported_background_fills'], 1)
        self.assertTrue(report['lossy'])

    def test_13_left_alignment_import(self):
        self._import_alignment(PP_ALIGN.LEFT, 'left')

    def test_14_center_alignment_import(self):
        self._import_alignment(PP_ALIGN.CENTER, 'center')

    def test_15_right_alignment_import(self):
        self._import_alignment(PP_ALIGN.RIGHT, 'right')

    def test_16_top_anchor_import(self):
        self._import_anchor(MSO_ANCHOR.TOP, 'top')

    def test_17_middle_anchor_import(self):
        self._import_anchor(MSO_ANCHOR.MIDDLE, 'middle')

    def test_18_bottom_anchor_import(self):
        self._import_anchor(MSO_ANCHOR.BOTTOM, 'bottom')

    def test_19_left_alignment_export(self):
        self._export_alignment('left', PP_ALIGN.LEFT)

    def test_20_center_alignment_export(self):
        self._export_alignment('center', PP_ALIGN.CENTER)

    def test_21_right_alignment_export(self):
        self._export_alignment('right', PP_ALIGN.RIGHT)

    def test_22_top_anchor_export(self):
        self._export_anchor('top', MSO_ANCHOR.TOP)

    def test_23_middle_anchor_export(self):
        self._export_anchor('middle', MSO_ANCHOR.MIDDLE)

    def test_24_bottom_anchor_export(self):
        self._export_anchor('bottom', MSO_ANCHOR.BOTTOM)

    def test_25_mixed_alignment_not_claimed_editable(self):
        deck, _, shape, path = self.source()
        shape.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
        shape.text_frame.paragraphs[1].alignment = PP_ALIGN.RIGHT
        deck.save(path)
        portable = pptx_to_portable(path)
        self.assertNotIn('textAlign', portable['activities'][0]['page_design']['widgets'][0]['style'])
        report = inspect_pptx_fidelity(path)
        self.assertEqual(report['mixed_paragraph_alignments'], 1)
        self.assertTrue(report['lossy'])

    def test_26_invalid_background_rejected_atomically(self):
        self._invalid('backgroundColor', 'red', 'background')

    def test_27_background_alpha_rejected_atomically(self):
        self._invalid('backgroundColor', '#11223344', 'background')

    def test_28_background_nonstring_rejected_atomically(self):
        self._invalid('backgroundColor', 123, 'background')

    def test_29_invalid_horizontal_alignment_rejected(self):
        self._invalid('textAlign', 'justify', 'text alignment')

    def test_30_invalid_vertical_alignment_rejected(self):
        self._invalid('verticalAlign', 'baseline', 'vertical alignment')

    def _import_background(self, rgb):
        deck, slide, _, path = self.source()
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(rgb)
        deck.save(path)
        self.assertEqual(pptx_to_portable(path)['activities'][0]['page_design']['backgroundColor'], '#' + rgb)

    def _export_background(self, color):
        destination = self.root / 'export.pptx'
        portable_to_pptx(definition(background=color), destination)
        self.assertEqual(str(Presentation(destination).slides[0].background.fill.fore_color.rgb), color[1:].upper())

    def _import_alignment(self, enum, expected):
        deck, _, shape, path = self.source()
        for paragraph in shape.text_frame.paragraphs:
            paragraph.alignment = enum
        deck.save(path)
        style = pptx_to_portable(path)['activities'][0]['page_design']['widgets'][0]['style']
        self.assertEqual(style['textAlign'], expected)
        self.assertEqual(inspect_pptx_fidelity(path)['editable_text_alignments'], 1)

    def _import_anchor(self, enum, expected):
        deck, _, shape, path = self.source()
        shape.text_frame.vertical_anchor = enum
        deck.save(path)
        style = pptx_to_portable(path)['activities'][0]['page_design']['widgets'][0]['style']
        self.assertEqual(style['verticalAlign'], expected)

    def _export_alignment(self, value, enum):
        destination = self.root / 'export.pptx'
        portable_to_pptx(definition(style={'textAlign': value}), destination)
        paragraphs = Presentation(destination).slides[0].shapes[0].text_frame.paragraphs
        self.assertTrue(all(p.alignment == enum for p in paragraphs))

    def _export_anchor(self, value, enum):
        destination = self.root / 'export.pptx'
        portable_to_pptx(definition(style={'verticalAlign': value}), destination)
        self.assertEqual(Presentation(destination).slides[0].shapes[0].text_frame.vertical_anchor, enum)

    def _invalid(self, key, value, error):
        data = definition()
        if key == 'backgroundColor':
            data['activities'][0]['page_design'][key] = value
        else:
            data['activities'][0]['page_design']['widgets'][0]['style'][key] = value
        destination = self.root / 'rejected.pptx'
        with self.assertRaisesRegex(ValueError, error):
            portable_to_pptx(data, destination)
        self.assertFalse(destination.exists())


if __name__ == '__main__':
    unittest.main()
