"""RC74: slide-master footer number must be at the canvas lower-right edge."""
import re
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]

class FooterCornerRC74(unittest.TestCase):
    def test_default_footer_occupies_bottom_strip_and_stays_dynamic(self):
        source = (ROOT / 'apps/web/src/slide-master.ts').read_text()
        self.assertIn('position:{x:0,y:94,width:100,height:6,rotation:0}', source)
        self.assertIn("footerShowPageNumber:true", source)
        self.assertIn("String(context.pageNumber)", source)

    def test_browser_layout_anchors_number_to_absolute_bottom_right_not_center(self):
        css = (ROOT / 'apps/web/src/style.css').read_text()
        declaration = re.search(r'\.projection-slide-footer \.slide-footer-number,\.canvas-slide-footer>span:last-child\{([^}]+)\}', css)
        self.assertIsNotNone(declaration)
        rules = declaration.group(1)
        for text in ('position:absolute;', 'right:0;', 'bottom:0;', 'top:auto;', 'transform:none;'):
            self.assertIn(text, rules)
        self.assertNotIn('translateY(', rules)
        self.assertIn('.audience-widget-layout>.page-widget[data-widget-kind="footer"]{padding:0!important;',css)
        main = (ROOT / 'apps/web/src/main.tsx').read_text()
        # Participant footer must stay *inside* the same 16:9 stage as other widgets.
        self.assertIn("w.kind==='footer'?<div className=\"projection-slide-footer\"", main)
        self.assertNotIn('participant-slide-footer" aria-label="Footer pagina attiva"', main)

    def test_powerpoint_number_at_exact_slide_right_and_bottom_without_text_insets(self):
        from pptx import Presentation
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from tools.pptx_portable import portable_to_pptx
        footer = {'id':'footer-x','kind':'footer','footerShowSessionTitle':True,
                  'footerShowPageNumber':True,'position':{'x':0,'y':94,'width':100,'height':6}}
        data = {'format':'quorivio-session','version':1,'session':{'title':'Sessione'},
                'activities':[{'title':'Prima','page_design':{'widgets':[footer]}},
                              {'title':'Seconda','page_design':{'widgets':[footer]}}]}
        with tempfile.TemporaryDirectory() as tmp:
            dest = Path(tmp)/'corner.pptx'
            portable_to_pptx(data,dest)
            pptx = Presentation(dest)
            for index,slide in enumerate(pptx.slides,1):
                number = next(shape for shape in slide.shapes if shape.has_text_frame and shape.text == str(index))
                self.assertAlmostEqual(number.left + number.width,pptx.slide_width,delta=3)
                self.assertAlmostEqual(number.top + number.height,pptx.slide_height,delta=3)
                self.assertEqual(number.text_frame.vertical_anchor,MSO_ANCHOR.BOTTOM)
                self.assertEqual(number.text_frame.margin_right,0)
                self.assertEqual(number.text_frame.margin_bottom,0)
                self.assertEqual(number.text_frame.paragraphs[0].alignment,PP_ALIGN.RIGHT)
                center = next(shape for shape in slide.shapes if shape.has_text_frame and shape.text=='Sessione')
                self.assertAlmostEqual(center.left + center.width/2,pptx.slide_width/2,delta=3)

if __name__ == '__main__':
    unittest.main()
