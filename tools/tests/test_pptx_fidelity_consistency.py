"""Regression coverage for the native PPTX importer and fidelity report parity."""
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'tools'))
from pptx_portable import _supported_image, inspect_pptx_fidelity, pptx_to_portable


class PptxFidelityConsistency(unittest.TestCase):
    def test_media_magic_must_match_declared_type(self):
        from PIL import Image
        import io
        buffer = io.BytesIO()
        Image.new('RGB', (1, 1), (255, 0, 0)).save(buffer, format='PNG')
        png = buffer.getvalue()
        buffer = io.BytesIO()
        Image.new('RGB', (1, 1), (255, 0, 0)).save(buffer, format='JPEG')
        jpg = buffer.getvalue()
        self.assertTrue(_supported_image(png, 'image/png'))
        self.assertTrue(_supported_image(jpg, 'image/jpeg'))
        self.assertFalse(_supported_image(jpg, 'image/png'))
        self.assertFalse(_supported_image(png, 'image/jpeg'))
        self.assertFalse(_supported_image(b'not an image', 'image/png'))
        self.assertFalse(_supported_image(b'', 'image/jpeg'))
        self.assertFalse(_supported_image(png, 'image/svg+xml'))

    def test_fidelity_matches_import_page_widget_limit(self):
        from pptx import Presentation
        from pptx.util import Inches
        with tempfile.TemporaryDirectory() as directory:
            filename = Path(directory) / 'many.pptx'
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            for index in range(18):
                slide.shapes.add_textbox(Inches(0.2), Inches(0.2), Inches(1), Inches(0.4)).text = f'Item {index}'
            deck.save(filename)
            converted = pptx_to_portable(filename)
            report = inspect_pptx_fidelity(filename)
            self.assertEqual(len(converted['activities'][0]['page_design']['widgets']), 16)
            self.assertEqual(report['editable_text_boxes'], 16)
            self.assertEqual(report['text_boxes_over_page_limit'], 2)
            self.assertTrue(report['lossy'])

if __name__ == '__main__':
    unittest.main()
