"""Security and roundtrip regression for the native editable PPTX path."""
import sys
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'tools'))
from pptx_portable import validate_pptx_archive, portable_to_pptx, pptx_to_portable, inspect_pptx_fidelity


def definition(widget):
    return {'format':'quorivio-session','version':1,'activities':[{'page_design':{'widgets':[widget]}}]}


class PptxAcceptanceHardening(unittest.TestCase):
    def test_requires_presentation_part(self):
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'fake.pptx'
            with ZipFile(path,'w') as z: z.writestr('[Content_Types].xml','<Types/>')
            with self.assertRaisesRegex(ValueError,'Not an OOXML'): validate_pptx_archive(path)

    def test_rejects_oversized_member(self):
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'bomb.pptx'
            with ZipFile(path,'w',compression=ZIP_DEFLATED) as z:
                z.writestr('[Content_Types].xml','<Types/>')
                z.writestr('ppt/presentation.xml','a'*(2*1024*1024))
            with self.assertRaisesRegex(ValueError,'compression ratio'): validate_pptx_archive(path)

    def test_rejects_symlink_archive_member(self):
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'symlink.pptx'
            with ZipFile(path,'w') as z:
                z.writestr('[Content_Types].xml','<Types/>')
                info=ZipInfo('ppt/asset');info.create_system=3;info.external_attr=(0o120777<<16)
                z.writestr(info,'../../secret')
            with self.assertRaisesRegex(ValueError,'symbolic links'): validate_pptx_archive(path)

    def test_rejects_external_relationships(self):
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'external.pptx'
            with ZipFile(path,'w') as z:
                z.writestr('[Content_Types].xml','<Types/>')
                z.writestr('ppt/presentation.xml','<p:presentation/>')
                z.writestr('ppt/_rels/presentation.xml.rels', '<Relationships><Relationship TargetMode="External" Target="https://example.org"/></Relationships>')
            with self.assertRaisesRegex(ValueError,'External OOXML'): validate_pptx_archive(path)

    def test_rejects_executable_js(self):
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'macro.pptx'
            with ZipFile(path,'w') as z:
                z.writestr('[Content_Types].xml','<Types/>')
                z.writestr('ppt/script.js','alert(1)')
            with self.assertRaisesRegex(ValueError,'Macro'): validate_pptx_archive(path)

    def test_rejects_zero_width_widget(self):
        with tempfile.TemporaryDirectory() as d:
            with self.assertRaisesRegex(ValueError,'geometry'):
                portable_to_pptx(definition({'kind':'text','label':'x','position':{'x':0,'y':0,'width':0,'height':20}}),Path(d)/'out.pptx')

    def test_invalid_uuid_raises_clear_error(self):
        with tempfile.TemporaryDirectory() as d:
            with self.assertRaisesRegex(ValueError,'asset ID'):
                portable_to_pptx(definition({'kind':'image','assetId':'not-a-uuid'}),Path(d)/'out.pptx',media_dir=d)

    def test_export_rejects_forged_image_bytes(self):
        import uuid
        with tempfile.TemporaryDirectory() as d:
            asset=str(uuid.uuid4());Path(d,asset).write_bytes(b'not a real png')
            with self.assertRaisesRegex(ValueError,'invalid image'):
                portable_to_pptx(definition({'kind':'image','assetId':asset}),Path(d)/'out.pptx',media_dir=d)

    def test_export_rejects_symlink_media_directory(self):
        import uuid
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);media=root/'real';media.mkdir();link=root/'link';link.symlink_to(media,target_is_directory=True)
            with self.assertRaisesRegex(ValueError,'image asset'):
                portable_to_pptx(definition({'kind':'image','assetId':str(uuid.uuid4())}),root/'out.pptx',media_dir=link)

    def test_long_imported_text_not_duplicated_on_export(self):
        from pptx import Presentation
        from pptx.util import Inches
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);src=root/'src.pptx';deck=Presentation()
            deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(2)).text='A'*280
            deck.save(src)
            data=pptx_to_portable(src)
            dest=root/'dest.pptx';portable_to_pptx(data,dest)
            self.assertEqual(Presentation(dest).slides[0].shapes[0].text,'A'*280)

    def test_fidelity_reports_truncated_text(self):
        from pptx import Presentation
        from pptx.util import Inches
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'long.pptx';deck=Presentation()
            deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(2)).text='x'*510
            deck.save(path)
            report=inspect_pptx_fidelity(path)
            self.assertEqual(report['truncated_text_boxes'],1)
            self.assertTrue(report['lossy'])

if __name__=='__main__':unittest.main()
