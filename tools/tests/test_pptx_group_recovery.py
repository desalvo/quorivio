"""End-to-end grouped PowerPoint shape recovery and fidelity reporting."""
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'tools'))
from pptx_portable import pptx_to_portable, inspect_pptx_fidelity


class GroupRecovery(unittest.TestCase):
    def test_grouped_editable_text_and_shape(self):
        from pptx import Presentation
        from pptx.util import Inches
        from pptx.enum.shapes import MSO_SHAPE
        with tempfile.TemporaryDirectory() as d:
            source = Path(d) / 'group.pptx'
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            group = slide.shapes.add_group_shape()
            textbox = group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
            textbox.text = 'Group heading'
            shape = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3), Inches(2), Inches(1), Inches(1))
            shape.text = 'Box'
            deck.save(source)
            imported = pptx_to_portable(source)
            widgets = imported['activities'][0]['page_design']['widgets']
            self.assertEqual([w['label'] for w in widgets], ['Group heading', 'Box'])
            self.assertLess(widgets[0]['position']['x'], widgets[1]['position']['x'])
            report = inspect_pptx_fidelity(source)
            self.assertEqual(report['editable_text_boxes'], 2)
            self.assertEqual(report['flattened_groups'], 1)
            self.assertTrue(report['lossy'])

    def test_nested_groups_and_page_limit_report(self):
        from pptx import Presentation
        from pptx.util import Inches
        from pptx.enum.shapes import MSO_SHAPE
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / 'nested.pptx'
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            outer = slide.shapes.add_group_shape()
            inner = outer.shapes.add_group_shape()
            for index in range(18):
                inner.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(index / 10 + 1),
                                       Inches(1), Inches(1), Inches(1)).text = f'Nested {index}'
            deck.save(path)
            imported = pptx_to_portable(path)
            report = inspect_pptx_fidelity(path)
            self.assertEqual(len(imported['activities'][0]['page_design']['widgets']), 16)
            self.assertEqual(report['flattened_groups'], 2)
            self.assertEqual(report['text_boxes_over_page_limit'], 2)

    def test_group_children_retain_scaled_positions(self):
        from pptx import Presentation
        from pptx.util import Inches
        from pptx.enum.shapes import MSO_SHAPE
        from pptx_portable import _flatten_shapes
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        group = slide.shapes.add_group_shape()
        group.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(1), Inches(1))
        group.shapes.add_shape(MSO_SHAPE.OVAL, Inches(2), Inches(1), Inches(1), Inches(1))
        group.width *= 2
        positions = list(_flatten_shapes(slide.shapes))
        self.assertEqual(len(positions), 2)
        self.assertEqual(positions[1].left - positions[0].left, Inches(2))

if __name__ == '__main__':
    unittest.main()
