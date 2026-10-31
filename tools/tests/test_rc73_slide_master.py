"""RC73: page title/footer master, numbered thumbnails and editable PPTX exports."""
import tempfile, unittest
from pathlib import Path
from pptx import Presentation
from tools.pptx_portable import portable_to_pptx
ROOT=Path(__file__).resolve().parents[2]
class SlideMasterRC73(unittest.TestCase):
 def test_editor_and_server_bind_the_defaults_and_allow_master_edit(self):
  ui=(ROOT/'apps/web/src/main.tsx').read_text()
  server=(ROOT/'apps/api/src/server.ts').read_text()
  for needle in ('newFooterWidget','footerShowSessionTitle','footerShowPageNumber','page-editor-thumbnail-caption','saveSlideMaster','masterTitleStyle'):
   self.assertIn(needle,ui)
  self.assertIn("'page-title',kind:'title'",server)
  self.assertIn("titleFollowPage:true",server)
  self.assertIn("'UPDATE activities SET page_design=$2::jsonb WHERE id=$1'",server)
  self.assertIn('titleStyle:slideMasterDraft',ui)
 def test_powerpoint_exports_title_in_master_style_and_live_footer_number(self):
  data={'format':'quorivio-session','version':1,'session':{'title':'Sessione uno','theme':{'master':{'titleStyle':{'fontSize':45,'bold':True,'fontFamily':'Georgia'}}}},'activities':[
   {'title':'Pagina A','page_design':{'widgets':[{'kind':'title','label':'Old title','titleFollowPage':True,'position':{'x':8,'y':3,'width':84,'height':12}},{'kind':'footer','label':'','footerShowSessionTitle':True,'footerShowPageNumber':True,'position':{'x':0,'y':94,'width':100,'height':6}}]}},
   {'title':'Pagina B','page_design':{'widgets':[{'kind':'title','label':'Old title','titleFollowPage':True,'position':{'x':8,'y':3,'width':84,'height':12}},{'kind':'footer','footerShowSessionTitle':True,'footerShowPageNumber':True,'position':{'x':0,'y':94,'width':100,'height':6}}]}}]}
  with tempfile.TemporaryDirectory() as folder:
   path=Path(folder)/'master.pptx';portable_to_pptx(data,path)
   pptx=Presentation(path)
   self.assertEqual(len(pptx.slides),2)
   first=[x.text for x in pptx.slides[0].shapes if x.has_text_frame]
   second=[x.text for x in pptx.slides[1].shapes if x.has_text_frame]
   self.assertIn('Pagina A',first);self.assertIn('Sessione uno',first);self.assertIn('1',first)
   self.assertIn('Pagina B',second);self.assertIn('2',second)
   title=next(x for x in pptx.slides[0].shapes if x.has_text_frame and x.text=='Pagina A')
   self.assertTrue(title.text_frame.paragraphs[0].runs[0].font.bold)
   self.assertEqual(title.text_frame.paragraphs[0].runs[0].font.name,'Georgia')
if __name__=='__main__':unittest.main()
