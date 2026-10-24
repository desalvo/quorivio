"""Recover meaningful PowerPoint chart titles without claiming editable chart data."""
import tempfile
import unittest
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from tools.pptx_portable import pptx_to_portable, inspect_pptx_fidelity

class PptxChartTitleTest(unittest.TestCase):
    def test_chart_title_is_editable_and_native_data_reported_lossy(self):
        with tempfile.TemporaryDirectory() as folder:
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            data = CategoryChartData()
            data.categories = ['A', 'B']
            data.add_series('Sales', (20, 30))
            chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(7), Inches(4), data).chart
            chart.has_title = True
            chart.chart_title.text_frame.text = 'Revenue by region'
            path = Path(folder) / 'chart.pptx'
            deck.save(path)
            widgets = pptx_to_portable(path)['activities'][0]['page_design']['widgets']
            report = inspect_pptx_fidelity(path)
            self.assertEqual([w['label'] for w in widgets], ['Revenue by region'])
            self.assertEqual(report['editable_chart_titles'], 1)
            self.assertEqual(report['unsupported_charts'], 1)
            self.assertTrue(report['lossy'])

if __name__ == '__main__': unittest.main()
