import json
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZipFile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'tools'))
from pptx_portable import portable_to_pptx, pptx_to_portable, inspect_pptx_fidelity

class PptxPortability(unittest.TestCase):
    def test_decorative_empty_shapes_roundtrip(self):
        from pptx import Presentation
        from pptx.util import Inches
        from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
        from pptx.dml.color import RGBColor
        with tempfile.TemporaryDirectory() as d:
            source = Path(d) / 'decorative.pptx'
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            for kind in (MSO_AUTO_SHAPE_TYPE.RECTANGLE, MSO_AUTO_SHAPE_TYPE.OVAL,
                         MSO_AUTO_SHAPE_TYPE.ISOSCELES_TRIANGLE):
                shape = slide.shapes.add_shape(kind, Inches(1), Inches(1), Inches(2), Inches(2))
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor(0x12, 0x34, 0x56)
            deck.save(source)
            converted = pptx_to_portable(source)
            widgets = converted['activities'][0]['page_design']['widgets']
            self.assertEqual(len(widgets), 3)
            self.assertTrue(all(w['kind'] == 'postit' and w['label'] == '' for w in widgets))
            self.assertEqual(inspect_pptx_fidelity(source)['unsupported_other_shapes'], 0)
            self.assertEqual(inspect_pptx_fidelity(source)['editable_decorative_shapes'], 3)
            target = Path(d) / 'out.pptx'
            portable_to_pptx(converted, target)
            exported = Presentation(target).slides[0].shapes
            self.assertEqual([sh.auto_shape_type for sh in exported],
                             [MSO_AUTO_SHAPE_TYPE.RECTANGLE, MSO_AUTO_SHAPE_TYPE.OVAL,
                              MSO_AUTO_SHAPE_TYPE.ISOSCELES_TRIANGLE])
            self.assertTrue(all(str(sh.fill.fore_color.rgb) == '123456' for sh in exported))

    def test_oval_and_triangle_roundtrip(self):
        from pptx import Presentation
        from pptx.util import Inches
        from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'shapes.pptx'
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            for kind, text in ((MSO_AUTO_SHAPE_TYPE.OVAL, 'Editable ellipse'),
                               (MSO_AUTO_SHAPE_TYPE.ISOSCELES_TRIANGLE, 'Editable triangle')):
                slide.shapes.add_shape(kind, Inches(1), Inches(1), Inches(2), Inches(2)).text = text
            deck.save(source)
            portable = pptx_to_portable(source)
            widgets = portable['activities'][0]['page_design']['widgets']
            self.assertEqual([w['style']['shapeType'] for w in widgets], ['ellipse', 'triangle'])
            self.assertTrue(all(w['kind'] == 'postit' for w in widgets))
            output = Path(directory) / 'result.pptx'
            portable_to_pptx(portable, output)
            exported = Presentation(output).slides[0].shapes
            self.assertEqual([shape.auto_shape_type for shape in exported],
                             [MSO_AUTO_SHAPE_TYPE.OVAL, MSO_AUTO_SHAPE_TYPE.ISOSCELES_TRIANGLE])


    def test_rounded_shape_colors_roundtrip(self):
        from pptx import Presentation
        from pptx.util import Inches
        from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
        from pptx.dml.color import RGBColor
        with tempfile.TemporaryDirectory() as d:
            source = Path(d) / 'colors.pptx'
            deck = Presentation()
            shape = deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(1), Inches(1), Inches(4), Inches(2))
            shape.text = 'Colored shape'
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(0x12, 0xAB, 0x34)
            shape.line.color.rgb = RGBColor(0x56, 0x78, 0x9A)
            deck.save(source)
            portable = pptx_to_portable(source)
            style = portable['activities'][0]['page_design']['widgets'][0]['style']
            self.assertEqual(style['fillColor'], '#12AB34')
            self.assertEqual(style['strokeColor'], '#56789A')
            self.assertTrue(style['rounded'])
            target = Path(d) / 'out.pptx'
            portable_to_pptx(portable, target)
            result = Presentation(target).slides[0].shapes[0]
            self.assertEqual(result.auto_shape_type, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE)
            self.assertEqual(str(result.fill.fore_color.rgb), '12AB34')
            self.assertEqual(str(result.line.color.rgb), '56789A')

    def test_rectangular_autoshape_remains_editable(self):
        from pptx import Presentation
        from pptx.util import Inches
        from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / 'shape.pptx'
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(1), Inches(1), Inches(3), Inches(1))
            rect.text = 'Editable rectangle'
            deck.save(path)
            result = pptx_to_portable(path)
            widget = result['activities'][0]['page_design']['widgets'][0]
            self.assertEqual(widget['kind'], 'postit')
            self.assertEqual(widget['label'], 'Editable rectangle')
            output = Path(d) / 'output.pptx'
            portable_to_pptx(result, output)
            exported = Presentation(output).slides[0].shapes[0]
            self.assertEqual(exported.shape_type, MSO_SHAPE_TYPE.AUTO_SHAPE)
            self.assertEqual(exported.text, 'Editable rectangle')

    def test_editable_text_roundtrip(self):
        from pptx import Presentation
        from pptx.util import Inches
        with tempfile.TemporaryDirectory() as d:
            f = Path(d) / 'initial.pptx'
            p = Presentation()
            slide = p.slides.add_slide(p.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2)).text = 'Editable heading'
            p.save(f)
            data = pptx_to_portable(f)
            self.assertEqual(data['activities'][0]['page_design']['widgets'][0]['label'], 'Editable heading')
            self.assertEqual(data['activities'][0]['type'], 'slide')
            out = Path(d) / 'roundtrip.pptx'
            portable_to_pptx(data, out)
            self.assertIn('Editable heading', Presentation(out).slides[0].shapes[0].text)
    def test_slide_title_and_rotation_survive_roundtrip(self):
        from pptx import Presentation
        from pptx.util import Inches
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'rotated.pptx'
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            heading = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
            heading.text = 'Project overview'
            heading.rotation = 15
            deck.save(source)
            portable = pptx_to_portable(source)
            self.assertEqual(portable['activities'][0]['title'], 'Project overview')
            self.assertEqual(portable['activities'][0]['page_design']['widgets'][0]['position']['rotation'], 15)
            target = Path(directory) / 'out.pptx'
            portable_to_pptx(portable, target)
            self.assertAlmostEqual(Presentation(target).slides[0].shapes[0].rotation, 15)

    def test_rejects_non_ooxml_and_macro_payloads(self):
        with tempfile.TemporaryDirectory() as d:
            invalid = Path(d) / 'invalid.pptx'
            with ZipFile(invalid, 'w') as archive:
                archive.writestr('dummy.txt', 'not a presentation')
            with self.assertRaisesRegex(ValueError, 'OOXML'):
                pptx_to_portable(invalid)
            with ZipFile(invalid, 'w') as archive:
                archive.writestr('[Content_Types].xml', '<Types/>')
                archive.writestr('ppt/vbaProject.bin', b'not executed')
            with self.assertRaisesRegex(ValueError, 'Macro'):
                pptx_to_portable(invalid)

    def test_unsupported_font_does_not_break_json_import(self):
        from pptx import Presentation
        from pptx.util import Inches
        with tempfile.TemporaryDirectory() as d:
            source = Path(d) / 'fonts.pptx'
            ppt = Presentation()
            shape = ppt.slides.add_slide(ppt.slide_layouts[6]).shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
            shape.text = 'Example'
            shape.text_frame.paragraphs[0].runs[0].font.name = 'Not-A-Quorivio-Font'
            ppt.save(source)
            widgets = pptx_to_portable(source)['activities'][0]['page_design']['widgets']
            self.assertNotIn('fontFamily', widgets[0]['style'])

    def test_underline_style_roundtrip(self):
        from pptx import Presentation
        from pptx.util import Inches
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / 'underline.pptx'
            deck = Presentation()
            shape = deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
            shape.text = 'Underlined title'
            shape.text_frame.paragraphs[0].runs[0].font.underline = True
            deck.save(path)
            portable = pptx_to_portable(path)
            self.assertTrue(portable['activities'][0]['page_design']['widgets'][0]['style']['underline'])
            target = Path(d) / 'export.pptx'
            portable_to_pptx(portable, target)
            self.assertTrue(Presentation(target).slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.underline)

    def test_loss_report_flags_pictures_without_serializing_assets(self):
        from pptx import Presentation
        from pptx.util import Inches
        from base64 import b64decode
        with tempfile.TemporaryDirectory() as d:
            image = Path(d) / 'one.png'
            image.write_bytes(b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAIAAACQd1PeAAAADElEQVR4nGP4z8AAAAMBAQDJ/pLvAAAAAElFTkSuQmCC'))
            path = Path(d) / 'with-image.pptx'
            deck = Presentation()
            deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_picture(str(image), Inches(1), Inches(1))
            deck.save(path)
            report = inspect_pptx_fidelity(path)
            self.assertEqual(report['editable_pictures'], 1)
            self.assertEqual(report['unsupported_pictures'], 0)
            self.assertNotIn('one.png', str(report))

    def test_unsupported_embedded_format_cannot_bypass_media_filter(self):
        from unittest.mock import patch
        from pptx import Presentation
        from pptx.util import Inches
        from base64 import b64decode
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            image = root / 'pixel.png'
            image.write_bytes(b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAIAAACQd1PeAAAADElEQVR4nGP4z8AAAAMBAQDJ/pLvAAAAAElFTkSuQmCC'))
            deck = Presentation()
            deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_picture(str(image), Inches(1), Inches(1))
            source = root / 'source.pptx'
            deck.save(source)
            # Simulate an unsupported embedded file type, keeping its bytes valid.
            from pptx.shapes.picture import Picture
            with patch.object(Picture, 'image', new_callable=lambda: property(
                lambda self: type('UnsupportedImage', (), {
                    'content_type': 'image/svg+xml', 'blob': b'<svg/>'})())):
                result = pptx_to_portable(source, media_dir=root / 'media')
            self.assertEqual(result['import_assets'], [])
            self.assertEqual(result['activities'][0]['page_design']['widgets'], [])

    def test_picture_extracts_to_editable_session_asset(self):
        from pptx import Presentation
        from pptx.util import Inches
        from base64 import b64decode
        with tempfile.TemporaryDirectory() as d:
            picture = Path(d) / 'picture.png'
            picture.write_bytes(b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAIAAACQd1PeAAAADElEQVR4nGP4z8AAAAMBAQDJ/pLvAAAAAElFTkSuQmCC'))
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_picture(str(picture), Inches(1), Inches(2), Inches(4))
            path = Path(d) / 'picture.pptx'
            deck.save(path)
            assets = Path(d) / 'media'
            result = pptx_to_portable(path, media_dir=assets)
            widget = result['activities'][0]['page_design']['widgets'][0]
            self.assertEqual(widget['kind'], 'image')
            self.assertEqual(widget['assetId'], result['import_assets'][0]['id'])
            self.assertEqual((assets / widget['assetId']).read_bytes(), picture.read_bytes())
            self.assertEqual(result['import_assets'][0]['bytes'], len(picture.read_bytes()))

    def test_picture_roundtrip_via_private_media_directory(self):
        from pptx import Presentation
        from pptx.util import Inches
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        from base64 import b64decode
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            picture = root / 'pixel.png'
            picture.write_bytes(b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAIAAACQd1PeAAAADElEQVR4nGP4z8AAAAMBAQDJ/pLvAAAAAElFTkSuQmCC'))
            deck = Presentation()
            shape = deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_picture(str(picture), Inches(1), Inches(1))
            shape.rotation = 25
            source = root / 'source.pptx'; deck.save(source)
            media = root / 'import-media'
            data = pptx_to_portable(source, media_dir=media)
            target = root / 'out.pptx'
            portable_to_pptx(data, target, media_dir=media)
            result = Presentation(target).slides[0].shapes[0]
            self.assertEqual(result.shape_type, MSO_SHAPE_TYPE.PICTURE)
            self.assertEqual(result.image.blob, picture.read_bytes())
            self.assertAlmostEqual(result.rotation, 25)
            with self.assertRaises(ValueError):
                portable_to_pptx(data, root / 'missing.pptx')

    def test_reused_picture_is_single_asset_with_multiple_widgets(self):
        from pptx import Presentation
        from pptx.util import Inches
        from base64 import b64decode
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            picture = root / 'pixel.png'
            picture.write_bytes(b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAIAAACQd1PeAAAADElEQVR4nGP4z8AAAAMBAQDJ/pLvAAAAAElFTkSuQmCC'))
            deck = Presentation()
            for _ in range(2):
                slide = deck.slides.add_slide(deck.slide_layouts[6])
                slide.shapes.add_picture(str(picture), Inches(1), Inches(1))
                slide.shapes.add_picture(str(picture), Inches(3), Inches(1))
            source = root / 'source.pptx'; deck.save(source)
            media = root / 'media'
            portable = pptx_to_portable(source, media_dir=media)
            self.assertEqual(len(portable['import_assets']), 1)
            ids = [w['assetId'] for a in portable['activities'] for w in a['page_design']['widgets']]
            self.assertEqual(len(ids), 4)
            self.assertEqual(len(set(ids)), 1)
            target = root / 'export.pptx'
            portable_to_pptx(portable, target, media_dir=media)
            self.assertEqual(sum(len(s.shapes) for s in Presentation(target).slides), 4)

    def test_rejects_unsafe_archive_member(self):
        with tempfile.TemporaryDirectory() as d:
            invalid = Path(d) / 'unsafe.pptx'
            with ZipFile(invalid, 'w') as archive:
                archive.writestr('[Content_Types].xml', '<Types/>')
                archive.writestr('../escape.txt', 'x')
            with self.assertRaisesRegex(ValueError, 'Unsafe'):
                pptx_to_portable(invalid)

    def test_native_small_logo_geometry_remains_editable(self):
        from pptx import Presentation
        from pptx.util import Inches
        with tempfile.TemporaryDirectory() as d:
            source = Path(d) / 'small-object.pptx'
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            shape = slide.shapes.add_textbox(Inches(7), Inches(3), Inches(.3), Inches(.2))
            shape.text = 'tiny'
            deck.save(source)
            portable = pptx_to_portable(source)
            pos = portable['activities'][0]['page_design']['widgets'][0]['position']
            self.assertLess(pos['width'], 12)
            self.assertLess(pos['height'], 8)
            target = Path(d) / 'result.pptx'
            portable_to_pptx(portable, target)
            result = Presentation(target).slides[0].shapes[0]
            self.assertAlmostEqual(result.width / Presentation(target).slide_width, shape.width / deck.slide_width, delta=.0001)
            self.assertAlmostEqual(result.height / Presentation(target).slide_height, shape.height / deck.slide_height, delta=.0001)

    def test_off_canvas_loss_is_reported(self):
        from pptx import Presentation
        from pptx.util import Inches
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / 'off-canvas.pptx'
            deck = Presentation()
            deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_textbox(-Inches(1), Inches(1), Inches(2), Inches(1)).text = 'clipped'
            deck.save(path)
            report = inspect_pptx_fidelity(path)
            self.assertEqual(report['off_canvas_shapes'], 1)
            self.assertTrue(report['lossy'])

    def test_invalid_geometry_rejected(self):
        with tempfile.TemporaryDirectory() as d:
            doc = {'format':'quorivio-session','version':1,'activities':[{'page_design':{'widgets':[{'kind':'text','position':{'x':90,'y':0,'width':30,'height':10}}]}}]}
            with self.assertRaises(ValueError):
                portable_to_pptx(doc, Path(d) / 'invalid.pptx')
    def test_validator_supports_slide(self):
        source = (ROOT / 'apps/api/src/server.ts').read_text()
        self.assertIn("['slide','poll','quiz','timed_quiz'].includes(a.type)", source)

if __name__ == '__main__':
    unittest.main()
