"""Test the isolated editable PowerPoint worker path and browser/API contract."""
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'tools'))
import conversion_worker

class BrowserPptxWorkerTests(unittest.TestCase):
    def test_worker_roundtrip_and_job_reservation(self):
        from pptx import Presentation
        from pptx.util import Inches
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            job = 'a' * 32
            workspace = root / job
            workspace.mkdir()
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = 'Editable demo'
            deck.save(workspace / 'source.pptx')
            with patch.object(conversion_worker, 'ROOT', root):
                imported = conversion_worker.execute_portable_job(job, 'import')
                self.assertEqual(imported['pages'], 1)
                payload = json.loads((workspace / 'portable.json').read_text())
                self.assertEqual(payload['activities'][0]['page_design']['widgets'][0]['label'], 'Editable demo')
                exported = conversion_worker.execute_portable_job(job, 'export')
                self.assertGreater(exported['bytes'], 0)
                self.assertEqual(Presentation(workspace / 'presentation.pptx').slides[0].shapes[0].text, 'Editable demo')
                self.assertFalse((workspace / '.conversion-lock').exists())

    def test_api_and_ui_routes_wired(self):
        api = (ROOT / 'apps/api/src/server.ts').read_text()
        web = (ROOT / 'apps/web/src/main.tsx').read_text()
        docker = (ROOT / 'Dockerfile.converter').read_text()
        self.assertIn("'/api/presentations/import-editable'", api)
        self.assertIn("'/api/sessions/:id/export.pptx'", api)
        self.assertIn("'/portable-import'", (ROOT / 'tools/conversion_worker.py').read_text())
        self.assertIn('importEditablePowerPoint', web)
        self.assertIn('export.pptx', web)
        self.assertIn('python-pptx==1.0.2', docker)
